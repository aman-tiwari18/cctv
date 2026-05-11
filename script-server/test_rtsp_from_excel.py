"""
RTSP Link Validation Script from Excel
Reads camera details from Excel and tests RTSP accessibility by capturing frames.
Sequential processing with colorful logging.

FIXES APPLIED:
  1. Token-first auth check (no longer rejected by unexpected "code" values)
  2. Full debug logging of server responses to diagnose auth issues
  3. TOKEN_EXPIRY_SECONDS raised to 3600 (1 hour)
  4. Dead-server tracking — servers that fail auth are skipped for remaining cameras
  5. Graceful CSV write even on partial results (KeyboardInterrupt safe)
  6. FFmpeg: replaced deprecated -stimeout with -timeout for RTSP (fixes "Unrecognized option" error)
  7. OpenCV: use a background thread with strict wall-clock timeout instead of
     relying on CAP_PROP_OPEN_TIMEOUT_MSEC (which many builds silently ignore)
  8. Frame save directory printed clearly in logs
"""

import argparse
import cv2
import hashlib
import json
import logging
import os
import shutil
import subprocess
import sys
import threading
import time
import urllib.parse
import uuid
import requests
from datetime import datetime

# Try to import Excel reading libraries
try:
    import pandas as pd
    EXCEL_LIBRARY = 'pandas'
except ImportError:
    try:
        from openpyxl import load_workbook
        EXCEL_LIBRARY = 'openpyxl'
    except ImportError:
        EXCEL_LIBRARY = None

# Try to import colorama for colored output
try:
    from colorama import Fore, Back, Style, init
    init(autoreset=True)
    COLORAMA_AVAILABLE = True
except ImportError:
    COLORAMA_AVAILABLE = False
    class Fore:
        RED = GREEN = YELLOW = BLUE = MAGENTA = CYAN = WHITE = RESET = ""
    class Back:
        RED = GREEN = YELLOW = BLUE = MAGENTA = CYAN = WHITE = RESET = ""
    class Style:
        BRIGHT = DIM = NORMAL = RESET_ALL = ""


# ============================================================
# Colorful Logging Setup
# ============================================================
class ColoredFormatter(logging.Formatter):
    FORMATS = {
        logging.DEBUG:    Fore.CYAN   + "%(asctime)s - [DEBUG]    - %(message)s" + Style.RESET_ALL,
        logging.INFO:     Fore.GREEN  + "%(asctime)s - [INFO]     - %(message)s" + Style.RESET_ALL,
        logging.WARNING:  Fore.YELLOW + "%(asctime)s - [WARNING]  - %(message)s" + Style.RESET_ALL,
        logging.ERROR:    Fore.RED    + "%(asctime)s - [ERROR]    - %(message)s" + Style.RESET_ALL,
        logging.CRITICAL: Fore.RED + Back.WHITE + "%(asctime)s - [CRITICAL] - %(message)s" + Style.RESET_ALL,
    }

    def format(self, record):
        log_fmt = self.FORMATS.get(record.levelno, self.FORMATS[logging.INFO])
        formatter = logging.Formatter(log_fmt, datefmt='%Y-%m-%d %H:%M:%S')
        return formatter.format(record)


logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)
logger.handlers.clear()

console_handler = logging.StreamHandler(sys.stdout)
console_handler.setLevel(logging.INFO)
console_handler.setFormatter(ColoredFormatter())
logger.addHandler(console_handler)

log_dir  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scripts", "logs", "rtsp_tests")
os.makedirs(log_dir, exist_ok=True)
log_file = os.path.join(log_dir, f"rtsp_test_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log")
file_handler = logging.FileHandler(log_file, mode='w', encoding='utf-8')
file_handler.setLevel(logging.DEBUG)
file_handler.setFormatter(logging.Formatter(
    '%(asctime)s - [%(levelname)s] - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
))
logger.addHandler(file_handler)


# ============================================================
# Configuration
# ============================================================
TIMEOUT_API          = 5    # seconds for HTTP API calls
TIMEOUT_OPENCV_HARD  = 6    # FIX #7: wall-clock thread-join limit for OpenCV (seconds)
TIMEOUT_FFMPEG       = 6    # seconds for FFmpeg subprocess

AUTH_USERNAME        = "system"
AUTH_PASSWORD        = "Coreiseven@@"
TOKEN_EXPIRY_SECONDS = 3600   # FIX #3: was 90s, now 1 hour
AUTH_MAX_ATTEMPTS    = 2

# Where captured frames are stored (absolute path, shown in logs)
FRAMES_BASE_DIR = os.path.join(os.getcwd(), "frames", "test_excel")


# ============================================================
# Authentication State
# ============================================================
server_tokens           = {}
server_token_timestamps = {}
dead_servers            = set()   # FIX #4: permanently failed servers
token_lock              = threading.Lock()


# ============================================================
# Colored Log Helpers
# ============================================================
def log_step(step_num, total_steps, message):
    logger.info(f"{Fore.CYAN}[Step {step_num}/{total_steps}]{Style.RESET_ALL} {message}")

def log_success(message):
    logger.info(f"{Fore.GREEN}✓ {message}{Style.RESET_ALL}")

def log_error(message):
    logger.error(f"{Fore.RED}✗ {message}{Style.RESET_ALL}")

def log_warning(message):
    logger.warning(f"{Fore.YELLOW}⚠ {message}{Style.RESET_ALL}")

def log_camera_header(cam_num, total_cams, camera_info):
    sep = "=" * 80
    logger.info(f"\n{Fore.MAGENTA}{sep}{Style.RESET_ALL}")
    logger.info(f"{Fore.MAGENTA}Camera {cam_num}/{total_cams}{Style.RESET_ALL}")
    logger.info(f"{Fore.MAGENTA}{sep}{Style.RESET_ALL}")
    logger.info(f"{Fore.CYAN}Row:{Style.RESET_ALL}       {camera_info['row']}")
    logger.info(f"{Fore.CYAN}Name:{Style.RESET_ALL}      {camera_info.get('camera_name', 'N/A')}")
    logger.info(f"{Fore.CYAN}Centre:{Style.RESET_ALL}    {camera_info.get('centre_name', 'N/A')}")
    logger.info(f"{Fore.CYAN}Server:{Style.RESET_ALL}    {camera_info['server_address']}")
    logger.info(f"{Fore.CYAN}Device ID:{Style.RESET_ALL} {camera_info['device_id']}")
    logger.info(f"{Fore.CYAN}Camera No:{Style.RESET_ALL} {camera_info['camera_no']}")
    logger.info(f"{Fore.MAGENTA}{sep}{Style.RESET_ALL}")


# ============================================================
# MD5 Signature
# ============================================================
def generate_signature(userName, password, realm, randomKey):
    temp = hashlib.md5(password.encode()).hexdigest()
    temp = hashlib.md5((userName + temp).encode()).hexdigest()
    temp = hashlib.md5(temp.encode()).hexdigest()
    temp = hashlib.md5((userName + ":" + realm + ":" + temp).encode()).hexdigest()
    return hashlib.md5((temp + ":" + randomKey).encode()).hexdigest()


# ============================================================
# HTTP Helper
# ============================================================
def make_api_call(url, payload, headers, timeout=None):
    if timeout is None:
        timeout = TIMEOUT_API
    try:
        response = requests.post(
            url, headers=headers, data=json.dumps(payload),
            timeout=timeout, verify=False
        )
        if response.status_code in (200, 401):
            return response
        logger.debug(f"Unexpected HTTP {response.status_code} from {url}")
        return None
    except requests.exceptions.Timeout:
        logger.debug(f"Timeout calling {url}")
        return None
    except Exception as e:
        logger.debug(f"API call exception: {e}")
        return None


def parse_json_response(response, context):
    if response is None:
        return None
    try:
        return response.json()
    except Exception as e:
        logger.debug(f"JSON parse error in '{context}': {e} | raw={response.text[:200]}")
        return None


# ============================================================
# Authentication  (FIX #1 + #2)
# ============================================================
def authenticate_server(server_address, max_attempts=AUTH_MAX_ATTEMPTS):
    """
    Authenticate with Camview server.
    FIX #1: Check for token presence first — don't reject on unexpected 'code'.
    FIX #2: Log full server responses at DEBUG level.
    """
    log_step(1, 3, f"Authenticating with server: {Fore.YELLOW}{server_address}{Style.RESET_ALL}")

    auth_url = f"https://{server_address}/admin/API/accounts/authorize"
    headers  = {"Content-Type": "application/json"}
    payload  = {"userName": AUTH_USERNAME, "ipAddress": "", "clientType": "WINPC_V1"}

    for attempt in range(1, max_attempts + 1):
        logger.info(f"  {Fore.CYAN}→{Style.RESET_ALL} Auth attempt {attempt}/{max_attempts}")

        # ── Step 1: challenge ─────────────────────────────────────────────
        res = make_api_call(auth_url, payload, headers)
        if res is None:
            log_warning(f"Step-1 API call failed (attempt {attempt})")
            time.sleep(attempt)
            continue

        resp_json = parse_json_response(res, "step-1 challenge")
        logger.debug(f"  Step-1 response: {resp_json}")

        if resp_json is None:
            log_warning(f"Step-1 non-JSON response (attempt {attempt})")
            time.sleep(attempt)
            continue

        realm     = resp_json.get("realm")
        randomKey = resp_json.get("randomKey")
        publicKey = resp_json.get("publickey") or resp_json.get("publicKey")
        cookie    = res.headers.get("Set-Cookie", "")

        if not all([realm, randomKey, publicKey]):
            log_warning(
                f"Step-1 missing fields (attempt {attempt}): "
                f"realm={realm!r} randomKey={randomKey!r} publicKey={publicKey!r}"
            )
            time.sleep(attempt)
            continue

        # ── Step 2: credential exchange ───────────────────────────────────
        logger.info(f"  {Fore.CYAN}→{Style.RESET_ALL} Generating signature and exchanging token...")
        sig = generate_signature(AUTH_USERNAME, AUTH_PASSWORD, realm, randomKey)
        payload2 = {
            "mac": "", "signature": sig, "userName": AUTH_USERNAME,
            "randomKey": randomKey, "publicKey": publicKey, "encryptType": "MD5",
            "ipAddress": "", "clientType": "WINPC_V1", "userType": "0"
        }
        headers2 = {"Content-Type": "application/json", "Cookie": cookie}
        res2 = make_api_call(auth_url, payload2, headers2)

        if res2 is None:
            log_warning(f"Step-2 API call failed (attempt {attempt})")
            time.sleep(attempt)
            continue

        resp_json2 = parse_json_response(res2, "step-2 token exchange")
        logger.debug(f"  Step-2 response: {resp_json2}")

        if resp_json2 is None:
            log_warning(f"Step-2 non-JSON response (attempt {attempt})")
            time.sleep(attempt)
            continue

        # FIX #1 — token-first: succeed if token present, regardless of 'code'
        token = resp_json2.get("token")
        if token:
            log_success(f"Authentication successful! Token: {token[:20]}...")
            return token

        code = resp_json2.get("code")
        log_warning(
            f"Step-2 returned no token (attempt {attempt}) — "
            f"code={code!r}, keys={list(resp_json2.keys())}"
        )
        time.sleep(attempt)

    log_error(f"Authentication failed after {max_attempts} attempts for {server_address}")
    return None


# ============================================================
# Token Cache  (FIX #3 expiry + FIX #4 dead-server)
# ============================================================
def get_server_token(server_address):
    with token_lock:
        if server_address in dead_servers:
            logger.debug(f"Skipping dead server: {server_address}")
            return None

        if server_address in server_tokens:
            elapsed = time.time() - server_token_timestamps.get(server_address, 0)
            if elapsed < TOKEN_EXPIRY_SECONDS:
                logger.debug(f"Using cached token for {server_address} (age {elapsed:.0f}s)")
                return server_tokens[server_address]

        token = authenticate_server(server_address)
        if token:
            server_tokens[server_address]           = token
            server_token_timestamps[server_address] = time.time()
        else:
            log_warning(f"Marking {server_address} as dead (all auth attempts failed)")
            dead_servers.add(server_address)
        return token


def add_token_to_url(rtsp_url, token):
    parsed    = urllib.parse.urlparse(rtsp_url)
    query     = urllib.parse.parse_qs(parsed.query)
    query['token'] = [token]
    new_query = urllib.parse.urlencode(query, doseq=True)
    return urllib.parse.urlunparse(parsed._replace(query=new_query))


# ============================================================
# Stream URL Fetching
# ============================================================
def start_video_stream(server_address, device_id, camera_no):
    log_step(2, 3, "Fetching RTSP URLs from server")

    token = get_server_token(server_address)
    if not token:
        log_error("No valid authentication token available")
        return None, None, False

    url     = f"https://{server_address}/admin/API/MTS/Video/StartVideo?token={token}"
    headers = {"Content-Type": "application/json", "X-Subject-Token": token}
    payload = {
        "clientType": "WINPC_V1", "project": "PSDK",
        "method": "MTS.Video.StartVideo",
        "data": {
            "streamType": "1",
            "optional": f"/admin/API/MTS/Video/StartVideo?token={token}",
            "channelId": f"{device_id}$1$0${camera_no}",
            "dataType": "2",
            "enableRtsps": "0"
        }
    }

    logger.info(f"  {Fore.CYAN}→{Style.RESET_ALL} Requesting stream for channel: {device_id}$1$0${camera_no}")
    res = make_api_call(url, payload, headers)

    if res is None:
        log_error("StartVideo API call failed")
        return None, None, False

    if res.status_code == 401:
        log_warning("Token expired (401) — clearing cache for this server")
        with token_lock:
            server_tokens.pop(server_address, None)
            server_token_timestamps.pop(server_address, None)
        return None, None, False

    if res.status_code != 200:
        log_error(f"StartVideo returned HTTP {res.status_code}")
        return None, None, False

    try:
        resp_json = res.json()
        logger.debug(f"StartVideo response: {resp_json}")
        data = resp_json.get("data")
        if not data:
            log_error("No 'data' in StartVideo response")
            return None, None, False

        urls = []
        for key in ("url", "url2"):
            raw = data.get(key)
            if not raw:
                continue
            parts = [p.strip() for p in raw.split('|') if p.strip()]
            urls.extend([urllib.parse.unquote(p) for p in parts])

        stream_token = data.get("token") or token

        if urls:
            log_success(f"Retrieved {len(urls)} RTSP URL(s)")
            for idx, u in enumerate(urls, 1):
                logger.info(f"  {Fore.CYAN}→{Style.RESET_ALL} URL {idx}: {u[:80]}...")
            return urls, stream_token, True

        log_error("No RTSP URLs found in StartVideo response")
        return None, None, False

    except Exception as e:
        log_error(f"Error parsing StartVideo response: {e}")
        return None, None, False


def build_rtsp_url(server_address, device_id, camera_no):
    try:
        if not server_address or not device_id or not camera_no:
            return None, False
        rtsp_urls, stream_token, success = start_video_stream(server_address, device_id, camera_no)
        if success and rtsp_urls and stream_token:
            return [add_token_to_url(u, stream_token) for u in rtsp_urls], True
        return None, False
    except Exception as e:
        logger.debug(f"build_rtsp_url exception: {e}")
        return None, False


# ============================================================
# FIX #7 — OpenCV with real wall-clock timeout
# ============================================================
def capture_frame_with_opencv(stream_url, id_str):
    """
    Capture a JPEG frame using OpenCV.

    CAP_PROP_OPEN_TIMEOUT_MSEC is silently ignored by many OpenCV/FFmpeg builds
    (your logs showed: "Stream timeout triggered after 30005 ms" despite 3000 ms set).
    We enforce the timeout via daemon thread + thread.join(TIMEOUT_OPENCV_HARD).
    """
    logger.info(
        f"  {Fore.CYAN}→{Style.RESET_ALL} OpenCV capture "
        f"(hard wall-clock timeout: {TIMEOUT_OPENCV_HARD}s)"
    )
    result = [None]

    def _do_capture():
        cap = None
        try:
            cap = cv2.VideoCapture(stream_url)
            # Hints only — not relied upon for actual cutoff
            cap.set(cv2.CAP_PROP_OPEN_TIMEOUT_MSEC, TIMEOUT_OPENCV_HARD * 1000)
            cap.set(cv2.CAP_PROP_READ_TIMEOUT_MSEC,  TIMEOUT_OPENCV_HARD * 1000)

            if not cap.isOpened():
                logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} OpenCV: could not open stream")
                return

            cap.read()           # skip first (often partial/corrupt) frame
            ret, frame = cap.read()

            if not ret or frame is None:
                logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} OpenCV: read() returned no frame")
                return

            save_dir = os.path.join(FRAMES_BASE_DIR, id_str)
            os.makedirs(save_dir, exist_ok=True)
            img_path = os.path.join(save_dir, f"frame_{uuid.uuid4().hex[:8]}.jpg")
            cv2.imwrite(img_path, frame)
            result[0] = img_path

        except Exception as e:
            logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} OpenCV exception: {e}")
        finally:
            try:
                if cap is not None:
                    cap.release()
            except Exception:
                pass

    thread = threading.Thread(target=_do_capture, daemon=True)
    thread.start()
    thread.join(timeout=TIMEOUT_OPENCV_HARD)

    if thread.is_alive():
        # Thread still blocked inside OpenCV — stream is unreachable or very slow
        logger.info(
            f"  {Fore.YELLOW}→{Style.RESET_ALL} OpenCV: hard timeout ({TIMEOUT_OPENCV_HARD}s) "
            f"— stream unreachable, moving on"
        )
        return None

    if result[0]:
        logger.info(f"  {Fore.GREEN}→{Style.RESET_ALL} OpenCV: frame saved → {result[0]}")
    return result[0]


# ============================================================
# FIX #6 — FFmpeg without deprecated -stimeout
# ============================================================
def capture_frame_with_ffmpeg(stream_url, id_str):
    """
    Capture a JPEG frame using FFmpeg.

    FIX #6: -stimeout was dropped in modern FFmpeg builds.
    Use -timeout (microseconds) instead — it works for RTSP, HTTP, and other
    libavformat demuxers.  Combine with -rtsp_transport tcp for reliability.
    """
    logger.info(
        f"  {Fore.CYAN}→{Style.RESET_ALL} FFmpeg capture "
        f"(timeout: {TIMEOUT_FFMPEG}s)"
    )

    if not shutil.which("ffmpeg"):
        logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} FFmpeg not found in PATH — skipping")
        return None

    save_dir  = os.path.join(FRAMES_BASE_DIR, id_str)
    os.makedirs(save_dir, exist_ok=True)
    img_path   = os.path.join(save_dir, f"frame_{uuid.uuid4().hex[:8]}.jpg")
    timeout_us = TIMEOUT_FFMPEG * 1_000_000  # FFmpeg expects microseconds

    cmd = ["ffmpeg", "-y"]

    if stream_url.startswith("rtsp"):
        cmd += [
            "-rtsp_transport", "tcp",      # force TCP (more reliable than UDP)
            "-timeout",        str(timeout_us),  # FIX #6: replaces -stimeout
        ]
    else:
        cmd += ["-timeout", str(timeout_us)]

    cmd += [
        "-i",        stream_url,
        "-frames:v", "1",
        "-q:v",      "2",          # JPEG quality (2=best … 31=worst)
        "-loglevel", "error",
        img_path,
    ]

    logger.debug(f"  FFmpeg cmd: {' '.join(cmd)}")

    try:
        proc = subprocess.run(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=TIMEOUT_FFMPEG + 3,   # hard kill margin for subprocess itself
        )

        stderr_msg = proc.stderr.decode(errors='replace').strip()

        if proc.returncode != 0 or not os.path.exists(img_path):
            short_err = stderr_msg[:300] if stderr_msg else "(no stderr output)"
            logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} FFmpeg failed: {short_err}")
            return None

        logger.info(f"  {Fore.GREEN}→{Style.RESET_ALL} FFmpeg: frame saved → {img_path}")
        return img_path

    except subprocess.TimeoutExpired:
        logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} FFmpeg: subprocess timeout ({TIMEOUT_FFMPEG}s)")
        return None
    except Exception as e:
        logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} FFmpeg exception: {e}")
        return None


def capture_frame_from_stream(stream_url, id_str):
    """Try OpenCV then FFmpeg. Returns (path, method) or (None, None)."""
    img_path = capture_frame_with_opencv(stream_url, id_str)
    if img_path:
        return img_path, "OpenCV"

    logger.info(f"  {Fore.YELLOW}→{Style.RESET_ALL} OpenCV failed — falling back to FFmpeg...")
    img_path = capture_frame_with_ffmpeg(stream_url, id_str)
    if img_path:
        return img_path, "FFmpeg"

    return None, None


# ============================================================
# Excel Reading
# ============================================================
def read_excel_pandas(excel_path):
    try:
        df = pd.read_excel(excel_path, engine='openpyxl')
        logger.info(f"Total rows in Excel: {len(df)}")
        logger.info(f"Original columns: {df.columns.tolist()}")

        df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
        logger.info(f"Normalised columns: {df.columns.tolist()}")

        def find_col(candidates):
            for col in df.columns:
                for c in candidates:
                    if c in col:
                        return col
            return None

        server_col = find_col(['cctv', 'server', 'ip', 'url'])
        device_col = find_col(['device'])
        camera_col = find_col(['camera_no', 'camera_num', 'cam_no', 'camera', 'cam'])
        name_col   = find_col(['camera_name', 'name'])
        centre_col = find_col(['centre', 'center'])

        logger.info(
            f"Detected → server:{server_col}  device:{device_col}  "
            f"camera:{camera_col}  name:{name_col}  centre:{centre_col}"
        )

        if not all([server_col, device_col, camera_col]):
            logger.error("Required columns not found in Excel!")
            return None

        cameras = []
        for idx, row in df.iterrows():
            cam = {
                'row':            idx + 2,
                'server_address': str(row.get(server_col) or '').strip(),
                'device_id':      str(row.get(device_col) or '').strip(),
                'camera_no':      str(row.get(camera_col) or '').strip(),
                'camera_name':    str(row.get(name_col)   or '').strip() if name_col   else '',
                'centre_name':    str(row.get(centre_col) or '').strip() if centre_col else '',
            }
            for key in ('server_address', 'device_id', 'camera_no'):
                if cam[key].lower() in ('nan', 'none', ''):
                    cam[key] = ''

            if cam['server_address'] and cam['device_id'] and cam['camera_no']:
                cameras.append(cam)
            else:
                logger.debug(f"Skipping row {idx+2}: missing fields → {cam}")

        log_success(f"Valid cameras parsed: {len(cameras)}")
        return cameras or None

    except Exception as e:
        logger.error(f"pandas read error: {e}", exc_info=True)
        return None


def read_excel_openpyxl(excel_path):
    try:
        wb = load_workbook(excel_path, read_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        logger.info(f"Excel columns: {headers}")

        col_map = {}
        for idx, header in enumerate(headers):
            if header is None:
                continue
            h = str(header).lower().strip()
            if any(k in h for k in ('cctv_url', 'server', 'ip')):
                col_map['server_address'] = idx
            elif 'device' in h:
                col_map['device_id'] = idx
            elif 'camera_no' in h or 'camera_num' in h:
                col_map['camera_no'] = idx
            elif 'camera_name' in h:
                col_map['camera_name'] = idx
            elif 'centre' in h or 'center' in h:
                col_map['centre_name'] = idx

        cameras = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            def _get(key, default_idx):
                i   = col_map.get(key, default_idx)
                val = row[i] if i < len(row) else None
                return str(val or '').strip()

            cam = {
                'row':            row_idx,
                'server_address': _get('server_address', 0),
                'device_id':      _get('device_id',      1),
                'camera_no':      _get('camera_no',      2),
                'camera_name':    _get('camera_name',    3),
                'centre_name':    _get('centre_name',    4),
            }
            if cam['server_address'] and cam['device_id'] and cam['camera_no']:
                cameras.append(cam)

        wb.close()
        return cameras or None

    except Exception as e:
        logger.error(f"openpyxl read error: {e}")
        return None


def read_excel_file(excel_path):
    if not os.path.exists(excel_path):
        logger.error(f"Excel file not found: {excel_path}")
        return None
    logger.info(f"Reading Excel file: {excel_path}")
    if EXCEL_LIBRARY == 'pandas':
        return read_excel_pandas(excel_path)
    elif EXCEL_LIBRARY == 'openpyxl':
        return read_excel_openpyxl(excel_path)
    else:
        logger.error("No Excel library found. Run:  pip install pandas openpyxl")
        return None


# ============================================================
# Camera Test
# ============================================================
def test_camera(cam, cam_num, total_cams):
    server_address = cam['server_address']
    device_id      = cam['device_id']
    camera_no      = cam['camera_no']
    id_str         = f"{device_id}_{camera_no}"

    log_camera_header(cam_num, total_cams, cam)

    result = {
        'row':            cam['row'],
        'server_address': server_address,
        'device_id':      device_id,
        'camera_no':      camera_no,
        'camera_name':    cam.get('camera_name', 'N/A'),
        'centre_name':    cam.get('centre_name', 'N/A'),
        'status':         'FAILED',
        'error':          None,
        'frame_path':     None,
        'method':         None,
        'duration':       0,
    }

    start_time = time.time()

    try:
        stream_urls, success = build_rtsp_url(server_address, device_id, camera_no)
        if not success or not stream_urls:
            result['error'] = "Failed to get RTSP URLs from server"
            log_error(result['error'])
            result['duration'] = round(time.time() - start_time, 2)
            return result

        log_step(3, 3, f"Capturing frame from {len(stream_urls)} URL(s)")

        for url_idx, stream_url in enumerate(stream_urls, 1):
            logger.info(f"\n{Fore.CYAN}Trying URL {url_idx}/{len(stream_urls)}{Style.RESET_ALL}")
            img_path, method = capture_frame_from_stream(stream_url, id_str)
            if img_path:
                result.update(status='SUCCESS', frame_path=img_path, method=method)
                result['duration'] = round(time.time() - start_time, 2)
                logger.info(f"\n{Fore.GREEN}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
                log_success(f"Camera {cam_num}/{total_cams} SUCCESSFUL!")
                logger.info(f"{Fore.GREEN}Method:   {method}{Style.RESET_ALL}")
                logger.info(f"{Fore.GREEN}Frame:    {img_path}{Style.RESET_ALL}")
                logger.info(f"{Fore.GREEN}Duration: {result['duration']}s{Style.RESET_ALL}")
                logger.info(f"{Fore.GREEN}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}\n")
                return result
            else:
                log_warning(f"URL {url_idx}/{len(stream_urls)} failed to capture frame")

        result['error'] = f"Frame capture failed on all {len(stream_urls)} URL(s)"
        log_error(result['error'])

    except Exception as e:
        result['error'] = f"Exception: {e}"
        log_error(result['error'])

    result['duration'] = round(time.time() - start_time, 2)
    logger.info(f"\n{Fore.RED}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
    log_error(f"Camera {cam_num}/{total_cams} FAILED!")
    logger.info(f"{Fore.RED}Error:    {result['error']}{Style.RESET_ALL}")
    logger.info(f"{Fore.RED}Duration: {result['duration']}s{Style.RESET_ALL}")
    logger.info(f"{Fore.RED}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}\n")
    return result


# ============================================================
# CSV Writer  (FIX #5 — safe on interrupt)
# ============================================================
def write_csv(results, output_path):
    try:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(
                "Row,Server,Device_ID,Camera_No,Camera_Name,Centre_Name,"
                "Status,Error,Frame_Path,Method,Duration_Sec\n"
            )
            for r in sorted(results, key=lambda x: x['row']):
                error = (r['error'] or '').replace('"', "'")
                f.write(
                    f"{r['row']},{r['server_address']},{r['device_id']},{r['camera_no']},"
                    f"\"{r['camera_name']}\",\"{r['centre_name']}\",{r['status']},"
                    f"\"{error}\",\"{r['frame_path'] or ''}\","
                    f"{r['method'] or ''},{r['duration']}\n"
                )
        log_success(f"Results saved → {output_path}")
        return True
    except Exception as e:
        log_error(f"Failed to write CSV: {e}")
        return False


# ============================================================
# Main
# ============================================================
def main():
    parser = argparse.ArgumentParser(
        description='RTSP Link Validation from Excel — Sequential'
    )
    parser.add_argument(
        '--excel',
        default='/data/cctv_dev/cctv/script-server/RTSP Link (SR, CR, NR, NWR)_CTGD 2026.xlsx',
        help='Path to Excel file'
    )
    parser.add_argument('--output',    default='rtsp_test_results.csv',
                        help='Output CSV path (default: rtsp_test_results.csv)')
    parser.add_argument('--debug',     action='store_true',
                        help='Enable DEBUG logging (shows raw server responses)')
    parser.add_argument('--start-row', type=int, default=1,
                        help='1-based camera index to resume from (skip already-tested cameras)')
    args = parser.parse_args()

    if args.debug:
        logger.setLevel(logging.DEBUG)
        for h in logger.handlers:
            h.setLevel(logging.DEBUG)

    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    # ── Banner ──────────────────────────────────────────────────────────────
    logger.info(f"\n{Fore.CYAN}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
    logger.info(f"{Fore.CYAN}{Style.BRIGHT}  RTSP Link Validation — Sequential Processing{Style.RESET_ALL}")
    logger.info(f"{Fore.CYAN}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}Log file:      {log_file}{Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}Excel file:    {args.excel}{Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}Output CSV:    {os.path.abspath(args.output)}{Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}Frames dir:    {FRAMES_BASE_DIR}/<device_id>_<cam_no>/{Style.RESET_ALL}")
    if args.debug:
        logger.info(f"{Fore.CYAN}DEBUG mode:    ON (raw server responses will be logged){Style.RESET_ALL}")
    if COLORAMA_AVAILABLE:
        log_success("Colorful logging enabled")
    else:
        log_warning("Install 'colorama' for colors:  pip install colorama")
    logger.info(f"{Fore.CYAN}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}\n")

    # ── Read Excel ───────────────────────────────────────────────────────────
    cameras = read_excel_file(args.excel)
    if not cameras:
        log_error("No valid camera rows found — aborting.")
        return 1

    if args.start_row > 1:
        cameras = cameras[args.start_row - 1:]
        log_warning(f"Resuming from camera index {args.start_row} ({len(cameras)} remaining)")

    log_success(f"Cameras to test: {len(cameras)}")
    logger.info(f"{Fore.CYAN}{'='*80}{Style.RESET_ALL}\n")

    total_cams = len(cameras)
    results    = []
    start_time = time.time()

    try:
        for idx, cam in enumerate(cameras, 1):
            result = test_camera(cam, idx, total_cams)
            results.append(result)

            success_count = sum(1 for r in results if r['status'] == 'SUCCESS')
            failed_count  = len(results) - success_count
            elapsed       = time.time() - start_time
            avg_time      = elapsed / len(results)
            remaining     = (total_cams - len(results)) * avg_time

            logger.info(f"{Fore.CYAN}{'='*80}{Style.RESET_ALL}")
            logger.info(
                f"{Fore.CYAN}Progress: {len(results)}/{total_cams} "
                f"({len(results)/total_cams*100:.1f}%){Style.RESET_ALL}"
            )
            logger.info(
                f"{Fore.GREEN}✓ Success: {success_count}{Style.RESET_ALL}  |  "
                f"{Fore.RED}✗ Failed: {failed_count}{Style.RESET_ALL}  |  "
                f"{Fore.YELLOW}Dead servers: {len(dead_servers)}{Style.RESET_ALL}"
            )
            logger.info(
                f"{Fore.YELLOW}⏱  Elapsed: {int(elapsed//60)}m {int(elapsed%60)}s  |  "
                f"Remaining: ~{int(remaining//60)}m {int(remaining%60)}s{Style.RESET_ALL}"
            )
            logger.info(f"{Fore.CYAN}{'='*80}{Style.RESET_ALL}\n")

    except KeyboardInterrupt:
        log_warning("\nInterrupted by user — saving partial results...")

    # ── Summary ──────────────────────────────────────────────────────────────
    total_time    = time.time() - start_time
    success_count = sum(1 for r in results if r['status'] == 'SUCCESS')
    failed_count  = len(results) - success_count

    logger.info(f"\n{Fore.MAGENTA}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
    logger.info(f"{Fore.MAGENTA}{Style.BRIGHT}FINAL SUMMARY{Style.RESET_ALL}")
    logger.info(f"{Fore.MAGENTA}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
    logger.info(f"{Fore.CYAN}Cameras tested:  {len(results)} / {total_cams}{Style.RESET_ALL}")
    logger.info(f"{Fore.GREEN}✓ Successful:    {success_count} ({success_count/max(len(results),1)*100:.1f}%){Style.RESET_ALL}")
    logger.info(f"{Fore.RED}✗ Failed:        {failed_count} ({failed_count/max(len(results),1)*100:.1f}%){Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}Dead servers:    {len(dead_servers)}{Style.RESET_ALL}")
    if dead_servers:
        for ds in sorted(dead_servers):
            logger.info(f"{Fore.YELLOW}  • {ds}{Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}⏱  Total time:   {int(total_time//60)}m {int(total_time%60)}s{Style.RESET_ALL}")
    if results:
        logger.info(f"{Fore.YELLOW}⏱  Avg/camera:  {total_time/len(results):.2f}s{Style.RESET_ALL}")
    logger.info(f"{Fore.YELLOW}Frames dir:      {FRAMES_BASE_DIR}{Style.RESET_ALL}")
    logger.info(f"{Fore.MAGENTA}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}\n")

    write_csv(results, args.output)

    if failed_count > 0:
        logger.info(f"\n{Fore.RED}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
        logger.info(f"{Fore.RED}{Style.BRIGHT}FAILED CAMERAS{Style.RESET_ALL}")
        logger.info(f"{Fore.RED}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")
        for r in sorted(results, key=lambda x: x['row']):
            if r['status'] == 'FAILED':
                logger.info(
                    f"{Fore.RED}Row {r['row']:>5}: {r['device_id']}_{r['camera_no']}"
                    f"  —  {r['error']}{Style.RESET_ALL}"
                )
        logger.info(f"{Fore.RED}{Style.BRIGHT}{'='*80}{Style.RESET_ALL}")

    logger.info(f"\n{Fore.YELLOW}Log saved → {log_file}{Style.RESET_ALL}")
    return 0 if failed_count == 0 else 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        logger.warning(f"\n{Fore.YELLOW}Interrupted by user{Style.RESET_ALL}")
        sys.exit(1)
    except Exception as e:
        logger.error(f"{Fore.RED}Fatal error: {e}{Style.RESET_ALL}", exc_info=True)
        sys.exit(1)
