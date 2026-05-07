"""
CCTV Analytics Report Generation API
=====================================
POST /generate-report
Body: { "exam_id": 1, "date": "2026-04-05" }

Flow:
  1. exams table → link_type (HLS / RTSP)
  2. exam_shifts table → shift timings (dynamic, not hardcoded)
  3. camera_links_hls → camera list (exam_id + shift_id + date)
  4. Scan detection image folders → ch_code se match
  5. Excel report generate → OUTPUT_DIR mein save
"""

import os
import re
import sys
from datetime import datetime, time

import openpyxl
from openpyxl.styles import Font
from flask import Flask, request, jsonify
import pymysql

# ── Existing db_client from detection pipeline ──
sys.path.insert(0, "/data/cctv_dev/cctv/script-server/scripts/detection/utils")
from db_client import (
    DB_CONFIG,
    get_hls_cameras,
    get_rtsp_cameras,
)

app = Flask(__name__)

# ─────────────────────────────────────────────
# PATHS
# ─────────────────────────────────────────────
RESULTS_FOLDER = "/data/cctv_dev/cctv/script-server/scripts/detection/results/tcs_detections"
OUTPUT_DIR     = "/data/cctv_backup/cctv-analytics-system"
BASE_IMAGE_URL = "http://www.examscctvmonitoring.com/cctv/images/results/new_tcs_results"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ─────────────────────────────────────────────
# DETECTION TYPE MAPPING  (filename prefix → label)
# ─────────────────────────────────────────────
PREFIX_TO_DETECTION = {
    "crowd":      "Crowd",
    "EntryExit":  "Entry Exit",
    "phone":      "Phone",
    "violation":  "Proximity",
    "server":     "Server Room",
    "tampering":  "Camera Tampering",
    "Inversion":  "Camera Inversion",
    "lessthan5":  "Crowd Threshold",
    "zero":       "Zero Person",
    "blind":      "Blind Spot",
}

NEW_FORMAT_FOLDERS = [
    "crowd_detection_results",
    "entry_exit_detection_results",
    "phone_detection_results",
    "proximity_detection_results",
    "server_room_detection_results",
    "camera_tampering_detection_results",
    "camera_inversion_detection_results",
    "crowd_detection_threshold",
    "zero_person_detection",
    "blind_spot_detection_results",
]


# ═════════════════════════════════════════════
#  EXTRA DB HELPERS  (jo db_client mein nahi hain)
# ═════════════════════════════════════════════

def _get_conn():
    """db_client ka DB_CONFIG reuse karke connection banao."""
    cfg = {k: v for k, v in DB_CONFIG.items() if k != "cursorclass"}
    return pymysql.connect(**cfg, cursorclass=pymysql.cursors.DictCursor)


def fetch_exam_info(exam_id: int) -> dict | None:
    """exams table se link_type + basic info lo."""
    sql = """
        SELECT id, exam_name, exam_body, link_type,
               start_date, end_date, is_active
        FROM exams
        WHERE id = %s
        LIMIT 1
    """
    try:
        with _get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, (exam_id,))
                return cur.fetchone()
    except Exception as e:
        raise RuntimeError(f"DB error (fetch_exam_info): {e}")


def fetch_shifts(exam_id: int) -> list:
    """
    exam_shifts table se is exam ki saari active shifts lo.
    MySQL TIME column (timedelta) → Python time object.
    """
    sql = """
        SELECT id, shift_name, start_time, end_time
        FROM exam_shifts
        WHERE exam_id = %s AND is_active = 1
        ORDER BY start_time
    """
    try:
        with _get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, (exam_id,))
                rows = cur.fetchall()
    except Exception as e:
        raise RuntimeError(f"DB error (fetch_shifts): {e}")

    return [
        {
            "id":         r["id"],
            "shift_name": r["shift_name"],
            "start_time": _to_time(r["start_time"]),
            "end_time":   _to_time(r["end_time"]),
        }
        for r in rows
    ]


# ─────────────────────────────────────────────
# UTILS
# ─────────────────────────────────────────────

def _to_time(val) -> time:
    """MySQL timedelta / time → Python time."""
    if isinstance(val, time):
        return val
    total = int(val.total_seconds())
    return time(total // 3600, (total % 3600) // 60, total % 60)


def get_shift_for_time(image_time: time, shifts: list) -> dict | None:
    for s in shifts:
        if s["start_time"] <= image_time <= s["end_time"]:
            return s
    return None


def extract_datetime_from_filename(filename: str) -> datetime | None:
    m = re.search(r'(\d{8})_(\d{6})', filename)
    if m:
        try:
            return datetime.strptime(m.group(1) + m.group(2), "%Y%m%d%H%M%S")
        except Exception:
            return None
    return None


def extract_ch_code_from_filename(filename: str) -> str | None:
    """
    New format: prefix_chXXXX_YYYYMMDD_HHMMSS.jpg
    Returns lowercase ch_code e.g. "ch6334"
    """
    m = re.search(r'_(ch\d+)_\d{8}_\d{6}', os.path.basename(filename), re.IGNORECASE)
    return m.group(1).lower() if m else None


def get_detection_type(filename: str) -> str:
    name = os.path.basename(filename)
    for prefix, label in PREFIX_TO_DETECTION.items():
        if name.startswith(prefix + "_"):
            return label
    return "Unknown"


def build_ch_lookup(cameras: list) -> dict:
    """ch_code.lower() → camera row"""
    return {
        str(c["ch_code"]).strip().lower(): c
        for c in cameras
        if c.get("ch_code")
    }


def auto_width(ws, max_width=40):
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.value), default=0
        )
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, max_width)


# ═════════════════════════════════════════════
#  DETECTION SCANNER
# ═════════════════════════════════════════════

def collect_detections(cameras: list, shifts: list, target_date: str):
    """
    Scan detection folders for target_date.
    Match image ch_code → camera row via ch_lookup.
    Returns: (shift_counts dict, detections_list)
    """
    ch_lookup    = build_ch_lookup(cameras)
    shift_counts = {s["shift_name"]: {} for s in shifts}
    detections   = []
    matched = unmatched = 0
    unmatched_samples = []

    for folder in NEW_FORMAT_FOLDERS:
        folder_path = os.path.join(RESULTS_FOLDER, folder)
        if not os.path.exists(folder_path):
            continue

        print(f"  📂 {folder}")

        for root, _, files in os.walk(folder_path):
            if "false_positives" in root.lower():
                continue

            for fname in files:
                if not fname.lower().endswith((".jpg", ".jpeg", ".png")):
                    continue

                dt = extract_datetime_from_filename(fname)
                if not dt or dt.strftime("%Y-%m-%d") != target_date:
                    continue

                shift = get_shift_for_time(dt.time(), shifts)
                if not shift:
                    continue

                ch = extract_ch_code_from_filename(fname)
                if not ch:
                    unmatched += 1
                    unmatched_samples.append(f"parse_fail | {fname}")
                    continue

                cam = ch_lookup.get(ch)
                if not cam:
                    unmatched += 1
                    if len(unmatched_samples) < 10:
                        unmatched_samples.append(f"ch={ch} not in DB | {fname}")
                    continue

                matched += 1
                sname = shift["shift_name"]
                key   = cam["ch_code"]
                shift_counts[sname][key] = shift_counts[sname].get(key, 0) + 1

                detections.append([
                    sname,
                    cam["ch_code"],
                    cam.get("camera_name"),
                    cam.get("centre_name"),
                    cam.get("city"),
                    cam.get("state"),
                    cam.get("tc_code"),
                    get_detection_type(fname),
                    dt.strftime("%Y-%m-%d %H:%M:%S"),
                    f"{BASE_IMAGE_URL}/{folder}/{fname}",
                    "View",
                ])

    print(f"  ✅ Matched: {matched}  ❌ Unmatched: {unmatched}")
    for s in unmatched_samples[:10]:
        print("    →", s)

    return shift_counts, detections


# ═════════════════════════════════════════════
#  EXCEL BUILDERS
# ═════════════════════════════════════════════

def create_shift_sheets(wb, cameras, shift_counts, shifts):
    headers = [
        "City", "State", "Center", "TC Code", "Ch Code",
        "Camera Name", "Cam Type", "Classification",
        "Events (Excl. Rejected)", "Rejected Events", "Camera Status",
    ]
    for shift in shifts:
        ws = wb.create_sheet(shift["shift_name"])
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            ws.cell(row=1, column=c).font = Font(bold=True)

        for cam in cameras:
            ch    = cam.get("ch_code", "")
            count = shift_counts.get(shift["shift_name"], {}).get(ch, 0)
            
            # Camera is online if: is_active = 1 AND status = 'Accessible' (case-insensitive)
            is_active = cam.get("is_active") == 1
            status_value = str(cam.get("status", "")).strip().lower()
            is_accessible = status_value == "accessible"
            camera_status = "Online" if (is_active and is_accessible) else "Offline"
            
            ws.append([
                cam.get("city"),
                cam.get("state"),
                cam.get("centre_name"),
                cam.get("tc_code"),
                ch,
                cam.get("camera_name"),
                cam.get("camera_category"),
                cam.get("classification"),
                count,
                0,
                camera_status,
            ])
        auto_width(ws)


def create_detections_sheet(wb, detections):
    ws = wb.create_sheet("Detections List")
    headers = [
        "Shift", "Ch Code", "Camera Name", "Center Name",
        "City", "State", "TC Code",
        "Detection Type", "Detection Time", "Image URL", "View",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for row in detections:
        ws.append(row)
    auto_width(ws)


def create_analysis_sheet(wb, shift_counts, cameras, shifts):
    ws  = wb.create_sheet("Analysis")
    row = 1
    total_all = 0

    for shift in shifts:
        sname = shift["shift_name"]
        ws.cell(row=row, column=1, value=f"{sname} ANALYSIS").font = Font(bold=True)
        row += 1

        total_cams = len(cameras)
        
        # Online cameras: is_active = 1 AND status = 'Accessible' (case-insensitive)
        online_cams = sum(
            1 for c in cameras 
            if c.get("is_active") == 1 and str(c.get("status", "")).strip().lower() == "accessible"
        )
        
        total_det   = sum(shift_counts.get(sname, {}).values())
        total_all  += total_det
        online_rate = (online_cams / total_cams * 100) if total_cams else 0

        for label, val in [
            ("Total Cameras",    total_cams),
            ("Online Cameras",   online_cams),
            ("Offline Cameras",  total_cams - online_cams),
            ("Total Detections", total_det),
            ("Online Rate",      f"{online_rate:.2f}%"),
            ("Shift Start",      str(shift["start_time"])),
            ("Shift End",        str(shift["end_time"])),
        ]:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=val)
            row += 1

        row += 1
        ws.cell(
            row=row, column=1,
            value=f"Top 5 Most Active Cameras ({sname})"
        ).font = Font(bold=True)
        row += 1

        cam_map = {c["ch_code"]: c for c in cameras}
        for i, (ch, cnt) in enumerate(
            sorted(shift_counts.get(sname, {}).items(), key=lambda x: -x[1])[:5], 1
        ):
            cam = cam_map.get(ch, {})
            ws.cell(row=row, column=1, value=f"{i}. {ch} — {cam.get('camera_name', 'N/A')}")
            ws.cell(row=row, column=2, value=f"{cnt} detections")
            row += 1

        row += 2

    ws.cell(row=row, column=1, value="OVERALL SUMMARY").font = Font(bold=True)
    row += 1
    for label, val in [
        ("Total Detections (All Shifts)", total_all),
        ("Total Shifts",                  len(shifts)),
        ("Total Cameras",                 len(cameras)),
    ]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=val)
        row += 1

    auto_width(ws)


# ═════════════════════════════════════════════
#  FLASK ROUTES
# ═════════════════════════════════════════════

@app.route("/generate-report", methods=["POST"])
def generate_report():
    body = request.get_json(force=True)

    # ── Validate input ──
    exam_id     = body.get("exam_id")
    target_date = body.get("date")      # "YYYY-MM-DD"

    if not exam_id or not target_date:
        return jsonify({"error": "exam_id and date are required"}), 400

    try:
        exam_id = int(exam_id)
    except (TypeError, ValueError):
        return jsonify({"error": "exam_id must be an integer"}), 400

    try:
        datetime.strptime(target_date, "%Y-%m-%d")
    except ValueError:
        return jsonify({"error": "date must be YYYY-MM-DD"}), 400

    # ── Step 1: Exam info ──
    try:
        exam = fetch_exam_info(exam_id)
    except RuntimeError as e:
        return jsonify({"error": str(e)}), 500

    if not exam:
        return jsonify({"error": f"exam_id {exam_id} not found"}), 404

    link_type = (exam.get("link_type") or "").upper()
    print(f"\n📋 Exam : {exam['exam_name']}")
    print(f"   Type : {link_type}  |  Date: {target_date}")

    # ── Step 2: Dynamic shifts ──
    try:
        shifts = fetch_shifts(exam_id)
    except RuntimeError as e:
        return jsonify({"error": str(e)}), 500

    if not shifts:
        return jsonify({"error": f"No active shifts for exam_id={exam_id}"}), 404

    print(f"⏰ Shifts: {[s['shift_name'] for s in shifts]}")

    # ── Step 3: Camera list via db_client ──
    cameras = []

    if link_type == "HLS":
        # Har shift ke cameras alag ho sakte hain → sab merge, ch_code se deduplicate
        seen_ch = set()
        for shift in shifts:
            rows = get_hls_cameras(exam_id, shift["id"], target_date)
            for cam in rows:
                ch = str(cam.get("ch_code", "")).strip().lower()
                if ch and ch not in seen_ch:
                    seen_ch.add(ch)
                    cameras.append(cam)

    elif link_type == "RTSP":
        # RTSP mein shift/date filter nahi hai db_client mein
        cameras = get_rtsp_cameras(exam_id)

    else:
        return jsonify({"error": f"Unsupported link_type='{link_type}'"}), 400

    if not cameras:
        return jsonify({
            "error": (
                f"No cameras found for exam_id={exam_id}, "
                f"link_type={link_type}, date={target_date}"
            )
        }), 404

    print(f"📷 Cameras loaded: {len(cameras)}")
    
    # Debug: Check status values
    status_counts = {}
    active_counts = {"active": 0, "inactive": 0}
    
    for cam in cameras:
        # Check status field
        status_val = cam.get("status")
        if status_val is None:
            status_key = "NULL"
        else:
            status_key = str(status_val).strip().lower()
        status_counts[status_key] = status_counts.get(status_key, 0) + 1
        
        # Check is_active field
        if cam.get("is_active") == 1:
            active_counts["active"] += 1
        else:
            active_counts["inactive"] += 1
    
    print(f"📊 Status distribution: {status_counts}")
    print(f"📊 Active distribution: {active_counts}")
    
    # Debug: Sample camera data (first 3 cameras)
    print(f"🔍 Sample camera data (first 3):")
    for i, cam in enumerate(cameras[:3], 1):
        print(f"  {i}. ch_code={cam.get('ch_code')}, is_active={cam.get('is_active')}, status={repr(cam.get('status'))}")
    
    # Debug: Check online cameras (is_active=1 AND status='accessible')
    online_count = sum(
        1 for c in cameras 
        if c.get("is_active") == 1 and str(c.get("status", "")).strip().lower() == "accessible"
    )
    print(f"🟢 Online cameras (active + accessible): {online_count}/{len(cameras)}")
    
    if online_count == 0:
        print("⚠️  WARNING: No online cameras found!")
        print("   Possible reasons:")
        print("   1. status field is NULL or empty")
        print("   2. status field has different value (not 'Accessible')")
        print("   3. is_active is 0 for all cameras")
        print("   4. Vision classification script hasn't run yet to set status")

    # ── Step 4: Scan detections ──
    print("🔍 Scanning detection folders...")
    shift_counts, detections = collect_detections(cameras, shifts, target_date)

    # ── Step 5: Build Excel ──
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    create_shift_sheets(wb, cameras, shift_counts, shifts)
    create_detections_sheet(wb, detections)
    create_analysis_sheet(wb, shift_counts, cameras, shifts)

    output_file = os.path.join(OUTPUT_DIR, f"exam_{exam_id}_{target_date}.xlsx")
    wb.save(output_file)

    total_det = sum(sum(sc.values()) for sc in shift_counts.values())
    print(f"✅ Report saved: {output_file}")

    return jsonify({
        "success":          True,
        "exam_id":          exam_id,
        "exam_name":        exam.get("exam_name"),
        "link_type":        link_type,
        "date":             target_date,
        "shifts":           [s["shift_name"] for s in shifts],
        "total_cameras":    len(cameras),
        "total_detections": total_det,
        "output_file":      output_file,
    }), 200


@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "ok", "time": datetime.now().isoformat()}), 200


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)