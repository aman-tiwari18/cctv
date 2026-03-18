#!/usr/bin/env python3
"""
Financial Report Generator for February 2026 Exam
Generates comprehensive financial analysis based on daily camera reports
"""

import os
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

BASE_DIR = '/home/suvam/projects/cctv/cctv-analytics-system'
REPORTS_DIR = os.path.join(BASE_DIR, 'output/daily camera and financial report for feb 26')
OUTPUT_DIR = os.path.join(BASE_DIR, 'output')

# Exam configuration for Feb 2026
EXAM_CONFIG = {
    'name': 'Combined Graduate Level Examination (CGLE) 2026 Tier II',
    'dates': [
        '2026-02-11', '2026-02-12', '2026-02-13', 
        '2026-02-16', '2026-02-17', '2026-02-18', '2026-02-19'
    ],
    'shifts': [
        {'name': 'Shift1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
        {'name': 'Shift2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
        {'name': 'Shift3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
    ]
}

# Pricing rules (same as previous exams)
# - Shift > 2 hours: ₹22 per camera
# - Shift ≤ 2 hours (multiple shifts): ₹14 per camera
# - Single shift only: ₹22 per camera

def calculate_rate(duration, num_shifts_scheduled):
    """Calculate rate per camera based on pricing rules"""
    if num_shifts_scheduled == 1:
        return 22, "Single shift"
    elif duration > 2.0:
        return 22, f"{duration}hrs > 2hrs"
    else:
        return 14, f"{duration}hrs ≤ 2hrs"

def process_daily_report(date_str):
    """Process a single day's report"""
    excel_file = os.path.join(REPORTS_DIR, f"{date_str}.xlsx")
    
    if not os.path.exists(excel_file):
        return None
    
    shift_data = {}
    
    for shift_info in EXAM_CONFIG['shifts']:
        shift_name = shift_info['name']
        try:
            df = pd.read_excel(excel_file, sheet_name=shift_name)
            
            if 'Camera Status' in df.columns:
                online_count = (df['Camera Status'] == 'Online').sum()
                offline_count = (df['Camera Status'] == 'Offline').sum()
                
                detections = 0
                if 'Events(Excluding Rejected Events)' in df.columns:
                    detections = int(df['Events(Excluding Rejected Events)'].sum())
                
                shift_data[shift_name] = {
                    'online': online_count,
                    'offline': offline_count,
                    'total': len(df),
                    'detections': detections,
                    'duration': shift_info['duration'],
                    'start': shift_info['start'],
                    'end': shift_info['end']
                }
        except Exception as e:
            print(f"  ⚠️ Could not read {shift_name} for {date_str}: {e}")
    
    return shift_data

print("="*80)
print("FINANCIAL REPORT GENERATOR - FEBRUARY 2026 EXAM")
print("="*80)
print(f"Exam: {EXAM_CONFIG['name']}")
print(f"Dates: {EXAM_CONFIG['dates'][0]} to {EXAM_CONFIG['dates'][-1]}")
print("="*80)
print()

# Collect all data
all_data = []
total_exam_cost = 0

for date_str in EXAM_CONFIG['dates']:
    print(f"📊 Processing {date_str}...")
    
    shift_data = process_daily_report(date_str)
    
    if not shift_data:
        print(f"  ⚠️ No data found for {date_str}")
        continue
    
    # Calculate cost for this day
    scheduled_shifts = [s for s in shift_data.keys() if shift_data[s]['online'] > 0]
    num_shifts_scheduled = len(scheduled_shifts)
    
    if num_shifts_scheduled == 0:
        print(f"  📅 {date_str}: No active shifts (all cameras offline)")
        continue
    
    day_cost = 0
    shift_details = []
    
    for shift_name in scheduled_shifts:
        data = shift_data[shift_name]
        cameras = data['online']
        duration = data['duration']
        
        rate, reason = calculate_rate(duration, num_shifts_scheduled)
        shift_cost = cameras * rate
        day_cost += shift_cost
        
        shift_details.append({
            'shift': shift_name,
            'cameras': cameras,
            'detections': data['detections'],
            'rate': rate,
            'cost': shift_cost,
            'reason': reason,
            'calculation': f"{cameras} cams × ₹{rate} = ₹{shift_cost:,}"
        })
    
    # Store day data
    day_data = {
        'date': date_str,
        'num_shifts': num_shifts_scheduled,
        'total_cameras': sum([s['cameras'] for s in shift_details]),
        'total_detections': sum([s['detections'] for s in shift_details]),
        'day_cost': day_cost,
        'shifts': shift_details,
        'all_shift_data': shift_data
    }
    
    all_data.append(day_data)
    total_exam_cost += day_cost
    
    # Print summary
    shift_summary = ', '.join([f"{s['shift']}: {s['cameras']} cams" for s in shift_details])
    print(f"  ✓ {date_str}: {shift_summary} = ₹{day_cost:,}")

print()
print("="*80)
print(f"TOTAL COST: ₹{total_exam_cost:,}")
print(f"Total Days: {len(all_data)}")
print(f"Avg Cost/Day: ₹{total_exam_cost/len(all_data):,.2f}" if all_data else "N/A")
print("="*80)
print()

# Generate Excel Report
print("📝 Generating Financial Report Excel...")

wb = Workbook()
wb.remove(wb.active)

# --- Sheet 1: Executive Summary ---
ws1 = wb.create_sheet("Executive Summary")

# Header
ws1.merge_cells('A1:D1')
ws1['A1'] = 'CCTV MONITORING FINANCIAL REPORT'
ws1['A1'].font = Font(bold=True, size=16)
ws1['A1'].alignment = Alignment(horizontal='center')

ws1.merge_cells('A2:D2')
ws1['A2'] = EXAM_CONFIG['name']
ws1['A2'].font = Font(bold=True, size=12, color='0066CC')
ws1['A2'].alignment = Alignment(horizontal='center')

ws1.merge_cells('A3:D3')
ws1['A3'] = f'Generated: {datetime.now().strftime("%B %d, %Y %H:%M IST")}'
ws1['A3'].alignment = Alignment(horizontal='center')
ws1['A3'].font = Font(italic=True, size=10)

ws1.append([])
ws1.append([])

# Total cost box
ws1.append(['💰 TOTAL COST', f'₹{total_exam_cost:,}'])
ws1[f'A{ws1.max_row}'].font = Font(bold=True, size=14)
ws1[f'B{ws1.max_row}'].font = Font(bold=True, size=14, color='006600')
ws1.merge_cells(f'B{ws1.max_row}:D{ws1.max_row}')

ws1.append([])

# Summary stats
ws1.append(['Total Days with Monitoring:', len(all_data)])
ws1.append(['Date Range:', f"{EXAM_CONFIG['dates'][0]} to {EXAM_CONFIG['dates'][-1]}"])
ws1.append(['Average Cost per Day:', f"₹{total_exam_cost/len(all_data):,.2f}" if all_data else 'N/A'])
ws1.append(['Total Camera-Shifts:', sum([d['total_cameras'] for d in all_data])])
ws1.append(['Total Detections:', f"{sum([d['total_detections'] for d in all_data]):,}"])

ws1.append([])
ws1.append([])

# Pricing criteria
ws1.append(['PRICING CRITERIA'])
ws1[f'A{ws1.max_row}'].font = Font(bold=True, size=12, color='0066CC')
ws1.append(['• Shift > 2 hours:', '₹22 per camera per shift'])
ws1.append(['• Shift ≤ 2 hours (multiple shifts):', '₹14 per camera per shift'])
ws1.append(['• Single shift only:', '₹22 per camera per shift'])

ws1.append([])
ws1.append([])

# Daily breakdown summary
ws1.append(['DAILY COST BREAKDOWN'])
ws1[f'A{ws1.max_row}'].font = Font(bold=True, size=12, color='0066CC')
ws1.append(['Date', 'Shifts Active', 'Total Cameras', 'Total Detections', 'Daily Cost'])

header_row = ws1.max_row
for col in range(1, 6):
    cell = ws1.cell(row=header_row, column=col)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.alignment = Alignment(horizontal='center')

for day_data in all_data:
    ws1.append([
        day_data['date'],
        day_data['num_shifts'],
        day_data['total_cameras'],
        f"{day_data['total_detections']:,}",
        f"₹{day_data['day_cost']:,}"
    ])

# Column widths
ws1.column_dimensions['A'].width = 25
ws1.column_dimensions['B'].width = 20
ws1.column_dimensions['C'].width = 18
ws1.column_dimensions['D'].width = 18

# --- Sheet 2: Detailed Daily Analysis ---
ws2 = wb.create_sheet("Detailed Daily Analysis")

ws2.append(['DETAILED DAILY ANALYSIS - SHIFT-BY-SHIFT'])
ws2.merge_cells('A1:I1')
ws2['A1'].font = Font(bold=True, size=14)
ws2['A1'].alignment = Alignment(horizontal='center')

ws2.append([])

headers = ['Date', 'Shift', 'Start Time', 'End Time', 'Online Cameras', 'Detections', 'Rate/Camera', 'Shift Cost', 'Calculation']
ws2.append(headers)

header_row = ws2.max_row
for col in range(1, len(headers) + 1):
    cell = ws2.cell(row=header_row, column=col)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.alignment = Alignment(horizontal='center')

for day_data in all_data:
    for shift in day_data['shifts']:
        shift_info = next(s for s in EXAM_CONFIG['shifts'] if s['name'] == shift['shift'])
        ws2.append([
            day_data['date'],
            shift['shift'],
            shift_info['start'],
            shift_info['end'],
            shift['cameras'],
            f"{shift['detections']:,}",
            f"₹{shift['rate']}",
            f"₹{shift['cost']:,}",
            shift['calculation']
        ])
    
    # Add daily total row
    ws2.append([
        '', '', '', 'DAILY TOTAL →', 
        day_data['total_cameras'],
        f"{day_data['total_detections']:,}",
        '',
        f"₹{day_data['day_cost']:,}",
        ''
    ])
    last_row = ws2.max_row
    for col in range(4, 9):
        ws2.cell(row=last_row, column=col).font = Font(bold=True)
        ws2.cell(row=last_row, column=col).fill = PatternFill(start_color='E8F4F8', end_color='E8F4F8', fill_type='solid')
    
    ws2.append([])  # Blank row

# Column widths
ws2.column_dimensions['A'].width = 15
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 12
ws2.column_dimensions['D'].width = 12
ws2.column_dimensions['E'].width = 16
ws2.column_dimensions['F'].width = 14
ws2.column_dimensions['G'].width = 14
ws2.column_dimensions['H'].width = 14
ws2.column_dimensions['I'].width = 30

# --- Sheet 3: Camera Status Summary ---
ws3 = wb.create_sheet("Camera Status Summary")

ws3.append(['CAMERA STATUS SUMMARY BY DATE'])
ws3.merge_cells('A1:K1')
ws3['A1'].font = Font(bold=True, size=14)
ws3['A1'].alignment = Alignment(horizontal='center')

ws3.append([])

headers = ['Date', 'Shift1 Online', 'Shift1 Detections', 'Shift2 Online', 'Shift2 Detections', 
           'Shift3 Online', 'Shift3 Detections', 'Total Online', 'Total Detections', 'Online Rate', 'Cost']
ws3.append(headers)

header_row = ws3.max_row
for col in range(1, len(headers) + 1):
    cell = ws3.cell(row=header_row, column=col)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.alignment = Alignment(horizontal='center')

for day_data in all_data:
    all_shift_data = day_data['all_shift_data']
    
    s1 = all_shift_data.get('Shift1', {})
    s2 = all_shift_data.get('Shift2', {})
    s3 = all_shift_data.get('Shift3', {})
    
    total_online = day_data['total_cameras']
    total_detections = day_data['total_detections']
    
    # Calculate total cameras (not camera-shifts)
    total_cameras = s1.get('total', 0)  # Should be same for all shifts
    online_rate = (total_online / (total_cameras * day_data['num_shifts']) * 100) if total_cameras > 0 else 0
    
    ws3.append([
        day_data['date'],
        s1.get('online', 0),
        f"{s1.get('detections', 0):,}",
        s2.get('online', 0),
        f"{s2.get('detections', 0):,}",
        s3.get('online', 0),
        f"{s3.get('detections', 0):,}",
        total_online,
        f"{total_detections:,}",
        f"{online_rate:.1f}%",
        f"₹{day_data['day_cost']:,}"
    ])

# Column widths
for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
    ws3.column_dimensions[col].width = 14

# Save the workbook
output_file = os.path.join(OUTPUT_DIR, 'Financial_Report_Feb2026_Exam.xlsx')
wb.save(output_file)

print(f"✅ Financial report saved to: {output_file}")
print()

# Generate Markdown summary
print("📝 Generating Markdown Summary...")

md_content = f"""# Financial Report - {EXAM_CONFIG['name']}

## 📊 Executive Summary

**Generated:** {datetime.now().strftime("%B %d, %Y %H:%M IST")}  
**Analysis File:** `Financial_Report_Feb2026_Exam.xlsx`

---

## 💰 Total Cost: ₹{total_exam_cost:,}

### Key Metrics

| Metric | Value |
|--------|-------|
| **Total Days** | {len(all_data)} |
| **Date Range** | {EXAM_CONFIG['dates'][0]} to {EXAM_CONFIG['dates'][-1]} |
| **Avg Cost/Day** | ₹{total_exam_cost/len(all_data):,.2f} |
| **Total Camera-Shifts** | {sum([d['total_cameras'] for d in all_data]):,} |
| **Total Detections** | {sum([d['total_detections'] for d in all_data]):,} |

---

## 📋 Pricing Criteria Applied

| Scenario | Rate per Camera per Shift |
|----------|---------------------------|
| Shift duration > 2 hours | ₹22 |
| Shift duration ≤ 2 hours (with multiple shifts) | ₹14 |
| Only 1 shift scheduled for the day | ₹22 |

---

## 📈 Daily Breakdown

| Date | Shifts | Cameras | Detections | Daily Cost |
|------|--------|---------|------------|------------|
"""

for day_data in all_data:
    md_content += f"| {day_data['date']} | {day_data['num_shifts']} | {day_data['total_cameras']} | {day_data['total_detections']:,} | ₹{day_data['day_cost']:,} |\n"

md_content += f"""
**Total:** {len(all_data)} days | {sum([d['total_cameras'] for d in all_data]):,} camera-shifts | {sum([d['total_detections'] for d in all_data]):,} detections | **₹{total_exam_cost:,}**

---

## 📊 Detailed Analysis by Day

"""

for day_data in all_data:
    md_content += f"""### {day_data['date']}

**Total Cost:** ₹{day_data['day_cost']:,} | **Shifts:** {day_data['num_shifts']}

| Shift | Time | Cameras | Detections | Rate | Cost | Calculation |
|-------|------|---------|------------|------|------|-------------|
"""
    for shift in day_data['shifts']:
        shift_info = next(s for s in EXAM_CONFIG['shifts'] if s['name'] == shift['shift'])
        md_content += f"| {shift['shift']} | {shift_info['start']}-{shift_info['end']} | {shift['cameras']} | {shift['detections']:,} | ₹{shift['rate']} | ₹{shift['cost']:,} | {shift['calculation']} |\n"
    
    md_content += "\n"

md_content += f"""---

## 📁 Files Generated

- **Excel Report:** `output/Financial_Report_Feb2026_Exam.xlsx`
- **Summary:** This file

---

**Note:** All calculations follow the standard CCTV monitoring pricing criteria. Rates are applied based on shift duration and number of shifts scheduled per day.

**Generated by:** Financial Report Generator  
**Date:** {datetime.now().strftime("%B %d, %Y")}
"""

# Save markdown
md_file = os.path.join(OUTPUT_DIR, 'FINANCIAL_REPORT_FEB2026.md')
with open(md_file, 'w') as f:
    f.write(md_content)

print(f"✅ Markdown summary saved to: {md_file}")
print()
print("="*80)
print("✅ REPORT GENERATION COMPLETE!")
print("="*80)
print()
print(f"📊 Total Cost: ₹{total_exam_cost:,}")
print(f"📅 Days Monitored: {len(all_data)}")
print(f"💵 Average Cost/Day: ₹{total_exam_cost/len(all_data):,.2f}" if all_data else "N/A")
print()
print("📁 Generated Files:")
print(f"   • {output_file}")
print(f"   • {md_file}")
print()
