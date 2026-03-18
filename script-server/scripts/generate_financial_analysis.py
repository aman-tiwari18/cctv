#!/usr/bin/env python3
"""
Generate Financial Analysis for All Exam Reports
Calculates CCTV analysis costs based on:
- Shift > 2 hours: ₹22 per camera per shift
- Shift ≤ 2 hours: ₹14 per camera per shift  
- Only 1 shift scheduled for day: ₹22 per camera per shift
"""

import os
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

BASE_DIR = '/home/suvam/projects/cctv/cctv-analytics-system'
OUTPUT_BASE = os.path.join(BASE_DIR, 'output')

# Exam configurations with shift timings
EXAMS = {
    'sicpo_exam_9th_to_12th_dec': {
        'name': 'SICPO Exam',
        'dates': ['2025-12-09', '2025-12-10', '2025-12-11', '2025-12-12'],
        'shifts': [
            {'name': 'Shift 1', 'start': '10:00', 'end': '12:00', 'duration': 2.0},
            {'name': 'Shift 2', 'start': '15:00', 'end': '17:00', 'duration': 2.0},
        ]
    },
    'ssc_constable_driver_16_17_dec': {
        'name': 'SSC Constable Driver',
        'dates': ['2025-12-16', '2025-12-17'],
        'shifts': [
            {'name': 'Shift 1', 'start': '10:00', 'end': '11:30', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '15:00', 'end': '16:30', 'duration': 1.5},
        ]
    },
    'ssc_constable_exec_18dec_6jan': {
        'name': 'SSC Constable Executive',
        'dates': [
            '2025-12-18', '2025-12-19', '2025-12-20', '2025-12-21', '2025-12-22',
            '2025-12-23', '2025-12-24', '2025-12-25', '2025-12-26', '2025-12-27',
            '2025-12-28', '2025-12-29', '2025-12-30', '2025-12-31', '2026-01-01',
            '2026-01-02', '2026-01-03', '2026-01-04', '2026-01-05', '2026-01-06'
        ],
        'shifts': [
            {'name': 'Shift 1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
            {'name': 'Shift 3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
        ]
    },
    'ssc_head_constable_min_7_12jan': {
        'name': 'SSC Head Constable Ministerial',
        'dates': ['2026-01-07', '2026-01-08', '2026-01-09', '2026-01-10', '2026-01-11', '2026-01-12'],
        'shifts': [
            {'name': 'Shift 1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
            {'name': 'Shift 3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
        ]
    },
    'head_constable_awo_16_21jan': {
        'name': 'Head Constable AWO/TPO',
        'dates': ['2026-01-16', '2026-01-17', '2026-01-18', '2026-01-19', '2026-01-20', '2026-01-21'],
        'shifts': [
            {'name': 'Shift 1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
            {'name': 'Shift 3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
        ]
    }
}

print("="*80)
print("FINANCIAL ANALYSIS GENERATOR - CCTV EXAM MONITORING")
print("="*80)

# Data collection
all_data = []
exam_summaries = {}

for exam_folder, config in EXAMS.items():
    exam_path = os.path.join(OUTPUT_BASE, exam_folder)
    
    if not os.path.exists(exam_path):
        print(f"⚠️  Skipping {config['name']} - folder not found")
        continue
    
    print(f"\n📊 Processing {config['name']}...")
    
    exam_data = []
    total_exam_cost = 0
    total_exam_cameras = 0
    
    for date_str in config['dates']:
        excel_file = os.path.join(exam_path, f"{date_str}.xlsx")
        
        if not os.path.exists(excel_file):
            print(f"  ⚠️  File not found: {date_str}.xlsx")
            continue
        
        try:
            # Use pandas for faster reading - just need shift sheets
            scheduled_shifts = []
            shift_cameras = {}
            
            for shift_info in config['shifts']:
                shift_name = shift_info['name']
                try:
                    # Read shift sheet with pandas (faster)
                    df = pd.read_excel(excel_file, sheet_name=shift_name)
                    
                    # Count online cameras
                    if 'Camera Status' in df.columns:
                        online_count = (df['Camera Status'] == 'Online').sum()
                        
                        if online_count > 0:
                            scheduled_shifts.append(shift_info)
                            shift_cameras[shift_name] = online_count
                except:
                    # Sheet doesn't exist or error reading
                    pass
            
            # Determine pricing based on criteria
            num_shifts_scheduled = len(scheduled_shifts)
            
            if num_shifts_scheduled == 0:
                print(f"  📅 {date_str}: No active shifts")
                continue
            
            # Calculate cost per shift
            day_cost = 0
            shift_details = []
            
            for shift_info in scheduled_shifts:
                shift_name = shift_info['name']
                duration = shift_info['duration']
                cameras = shift_cameras.get(shift_name, 0)
                
                # Apply pricing rules
                if num_shifts_scheduled == 1:
                    # Only 1 shift scheduled → ₹22
                    rate = 22
                elif duration > 2.0:
                    # Shift > 2 hours → ₹22
                    rate = 22
                else:
                    # Shift ≤ 2 hours → ₹14
                    rate = 14
                
                shift_cost = cameras * rate
                day_cost += shift_cost
                
                shift_details.append({
                    'shift': shift_name,
                    'start': shift_info['start'],
                    'end': shift_info['end'],
                    'duration': duration,
                    'cameras': cameras,
                    'rate': rate,
                    'cost': shift_cost
                })
            
            # Store data
            day_data = {
                'Exam': config['name'],
                'Date': date_str,
                'Num Shifts Scheduled': num_shifts_scheduled,
                'Total Cameras (all shifts)': sum(shift_cameras.values()),
                'Day Cost': day_cost,
                'Shifts': shift_details
            }
            
            exam_data.append(day_data)
            all_data.append(day_data)
            total_exam_cost += day_cost
            
            print(f"  ✓ {date_str}: {num_shifts_scheduled} shifts, ₹{day_cost:,}")
            
        except Exception as e:
            print(f"  ❌ Error processing {date_str}: {e}")
    
    exam_summaries[config['name']] = {
        'days': len(exam_data),
        'total_cost': total_exam_cost,
        'data': exam_data
    }
    
    print(f"  📈 Total for {config['name']}: ₹{total_exam_cost:,}")

# Create Financial Analysis Excel
print("\n📝 Generating Financial Analysis Excel...")

output_file = os.path.join(OUTPUT_BASE, 'Financial_Analysis_CCTV_Monitoring.xlsx')
wb = Workbook()

# Remove default sheet
if 'Sheet' in wb.sheetnames:
    wb.remove(wb['Sheet'])

# ============================================================================
# Sheet 1: Executive Summary
# ============================================================================
ws_summary = wb.create_sheet('Executive Summary', 0)
ws_summary.column_dimensions['A'].width = 35
ws_summary.column_dimensions['B'].width = 20
ws_summary.column_dimensions['C'].width = 20
ws_summary.column_dimensions['D'].width = 20

# Header styling
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=12)

row = 1
ws_summary.cell(row, 1, "CCTV MONITORING FINANCIAL ANALYSIS").font = Font(bold=True, size=16, color="1F4E78")
ws_summary.merge_cells(f'A{row}:D{row}')
row += 1

ws_summary.cell(row, 1, f"Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}").font = Font(size=10, italic=True)
ws_summary.merge_cells(f'A{row}:D{row}')
row += 2

# Pricing Criteria
ws_summary.cell(row, 1, "PRICING CRITERIA").font = Font(bold=True, size=12, color="1F4E78")
row += 1

ws_summary.cell(row, 1, "• Shift > 2 hours:")
ws_summary.cell(row, 2, "₹22 per camera per shift")
row += 1
ws_summary.cell(row, 1, "• Shift ≤ 2 hours:")
ws_summary.cell(row, 2, "₹14 per camera per shift")
row += 1
ws_summary.cell(row, 1, "• Only 1 shift scheduled:")
ws_summary.cell(row, 2, "₹22 per camera per shift")
row += 2

# Overall Summary
ws_summary.cell(row, 1, "OVERALL SUMMARY").font = Font(bold=True, size=12, color="1F4E78")
row += 1

total_all_cost = sum(s['total_cost'] for s in exam_summaries.values())
total_all_days = sum(s['days'] for s in exam_summaries.values())

ws_summary.cell(row, 1, "Total Exams Analyzed:")
ws_summary.cell(row, 2, len(exam_summaries)).font = Font(bold=True)
row += 1
ws_summary.cell(row, 1, "Total Days Monitored:")
ws_summary.cell(row, 2, total_all_days).font = Font(bold=True)
row += 1
ws_summary.cell(row, 1, "Total Cost:")
ws_summary.cell(row, 2, f"₹{total_all_cost:,}").font = Font(bold=True, size=12, color="00B050")
row += 2

# Exam-wise Summary
ws_summary.cell(row, 1, "EXAM-WISE BREAKDOWN").font = Font(bold=True, size=12, color="1F4E78")
row += 1

ws_summary.cell(row, 1, "Exam Name").fill = header_fill
ws_summary.cell(row, 1).font = header_font
ws_summary.cell(row, 2, "Days").fill = header_fill
ws_summary.cell(row, 2).font = header_font
ws_summary.cell(row, 3, "Total Cost").fill = header_fill
ws_summary.cell(row, 3).font = header_font
ws_summary.cell(row, 4, "% of Total").fill = header_fill
ws_summary.cell(row, 4).font = header_font
row += 1

for exam_name, summary in exam_summaries.items():
    ws_summary.cell(row, 1, exam_name)
    ws_summary.cell(row, 2, summary['days'])
    ws_summary.cell(row, 3, f"₹{summary['total_cost']:,}")
    percentage = (summary['total_cost'] / total_all_cost * 100) if total_all_cost > 0 else 0
    ws_summary.cell(row, 4, f"{percentage:.1f}%")
    row += 1

# ============================================================================
# Sheet 2: Detailed Daily Analysis
# ============================================================================
ws_detail = wb.create_sheet('Daily Analysis', 1)
ws_detail.column_dimensions['A'].width = 30
ws_detail.column_dimensions['B'].width = 15
ws_detail.column_dimensions['C'].width = 15
ws_detail.column_dimensions['D'].width = 20
ws_detail.column_dimensions['E'].width = 15
ws_detail.column_dimensions['F'].width = 15
ws_detail.column_dimensions['G'].width = 15
ws_detail.column_dimensions['H'].width = 20

row = 1
ws_detail.cell(row, 1, "DAILY COST ANALYSIS").font = Font(bold=True, size=14, color="1F4E78")
ws_detail.merge_cells(f'A{row}:H{row}')
row += 2

# Header
headers = ['Exam', 'Date', 'Day of Week', 'Shifts Scheduled', 'Shift 1 Cameras', 'Shift 2 Cameras', 'Shift 3 Cameras', 'Daily Cost']
for col, header in enumerate(headers, 1):
    cell = ws_detail.cell(row, col, header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
row += 1

# Data rows
for day_data in all_data:
    date_obj = datetime.strptime(day_data['Date'], '%Y-%m-%d')
    day_name = date_obj.strftime('%A')
    
    ws_detail.cell(row, 1, day_data['Exam'])
    ws_detail.cell(row, 2, day_data['Date'])
    ws_detail.cell(row, 3, day_name)
    ws_detail.cell(row, 4, day_data['Num Shifts Scheduled'])
    
    # Shift cameras
    for col_idx, shift_detail in enumerate(day_data['Shifts'], 0):
        if col_idx < 3:  # Max 3 shifts
            ws_detail.cell(row, 5 + col_idx, shift_detail['cameras'])
    
    ws_detail.cell(row, 8, f"₹{day_data['Day Cost']:,}")
    row += 1

# ============================================================================
# Sheet 3: Shift-wise Details
# ============================================================================
ws_shifts = wb.create_sheet('Shift Details', 2)
ws_shifts.column_dimensions['A'].width = 30
ws_shifts.column_dimensions['B'].width = 15
ws_shifts.column_dimensions['C'].width = 12
ws_shifts.column_dimensions['D'].width = 12
ws_shifts.column_dimensions['E'].width = 12
ws_shifts.column_dimensions['F'].width = 15
ws_shifts.column_dimensions['G'].width = 12
ws_shifts.column_dimensions['H'].width = 15

row = 1
ws_shifts.cell(row, 1, "SHIFT-WISE COST BREAKDOWN").font = Font(bold=True, size=14, color="1F4E78")
ws_shifts.merge_cells(f'A{row}:H{row}')
row += 2

# Header
headers = ['Exam', 'Date', 'Shift', 'Start Time', 'End Time', 'Duration (hrs)', 'Cameras', 'Rate', 'Shift Cost']
for col, header in enumerate(headers, 1):
    cell = ws_shifts.cell(row, col, header)
    cell.fill = header_fill
    cell.font = header_font
row += 1

# Data rows
for day_data in all_data:
    for shift_detail in day_data['Shifts']:
        ws_shifts.cell(row, 1, day_data['Exam'])
        ws_shifts.cell(row, 2, day_data['Date'])
        ws_shifts.cell(row, 3, shift_detail['shift'])
        ws_shifts.cell(row, 4, shift_detail['start'])
        ws_shifts.cell(row, 5, shift_detail['end'])
        ws_shifts.cell(row, 6, shift_detail['duration'])
        ws_shifts.cell(row, 7, shift_detail['cameras'])
        ws_shifts.cell(row, 8, f"₹{shift_detail['rate']}")
        ws_shifts.cell(row, 9, f"₹{shift_detail['cost']:,}")
        row += 1

# ============================================================================
# Sheet 4: Exam-wise Summary
# ============================================================================
ws_exam = wb.create_sheet('Exam Summary', 3)
ws_exam.column_dimensions['A'].width = 35
ws_exam.column_dimensions['B'].width = 15
ws_exam.column_dimensions['C'].width = 15
ws_exam.column_dimensions['D'].width = 20
ws_exam.column_dimensions['E'].width = 20

row = 1
ws_exam.cell(row, 1, "EXAM-WISE FINANCIAL SUMMARY").font = Font(bold=True, size=14, color="1F4E78")
ws_exam.merge_cells(f'A{row}:E{row}')
row += 2

for exam_name, summary in exam_summaries.items():
    ws_exam.cell(row, 1, exam_name).font = Font(bold=True, size=12, color="1F4E78")
    row += 1
    
    ws_exam.cell(row, 1, "Total Days:")
    ws_exam.cell(row, 2, summary['days'])
    row += 1
    
    ws_exam.cell(row, 1, "Total Cost:")
    ws_exam.cell(row, 2, f"₹{summary['total_cost']:,}").font = Font(bold=True, color="00B050")
    row += 1
    
    ws_exam.cell(row, 1, "Average Cost per Day:")
    avg_cost = summary['total_cost'] / summary['days'] if summary['days'] > 0 else 0
    ws_exam.cell(row, 2, f"₹{avg_cost:,.2f}")
    row += 2

# Save workbook
wb.save(output_file)

print(f"\n✅ Financial Analysis generated: {output_file}")
print("\n" + "="*80)
print("SUMMARY")
print("="*80)
print(f"Total Exams: {len(exam_summaries)}")
print(f"Total Days: {total_all_days}")
print(f"Total Cost: ₹{total_all_cost:,}")
print("\nExam-wise breakdown:")
for exam_name, summary in exam_summaries.items():
    print(f"  • {exam_name}: ₹{summary['total_cost']:,} ({summary['days']} days)")
print("\n" + "="*80)
