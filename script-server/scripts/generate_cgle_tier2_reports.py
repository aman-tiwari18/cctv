#!/usr/bin/env python3
"""
Generate CGLE Tier 2 Exam Reports for Jan 19, 2026
Special case: 2 separate sheets for Paper I and Paper II
Creates reports with 2 shifts:
- Shift 1 (11:00-13:15) - Paper I
- Shift 2 (16:30-18:30) - Paper II
- Daily analysis
- Detections list with view links
- Proof of work folder with detection images
"""

import os
import re
import shutil
import pandas as pd
from datetime import datetime
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font

# Configuration
BASE_DIR = '/home/suvam/projects/cctv/cctv-analytics-system'
RESULTS_DIR = os.path.join(BASE_DIR, 'results')
OUTPUT_DIR = os.path.join(BASE_DIR, 'output/cgle_tier2_19jan')

# Two separate camera sheets for Paper I and Paper II
CAMERA_SHEET_PAPER1 = os.path.join(BASE_DIR, 'data/RTSP LINKS/CGLE TIER 2_19th_jan/CCTV Camera Report_ CGLE 2025 Tier II Paper I 19th Jan 2026.xlsx')
CAMERA_SHEET_PAPER2 = os.path.join(BASE_DIR, 'data/RTSP LINKS/CGLE TIER 2_19th_jan/CCTV Camera Report_CGLE 2025 Tier II Paper II Statistics 19th Jan 2026.xlsx')

IMAGE_BASE_URL = 'http://103.42.50.241/cctv/images/'

# Single date
DATES = ['20260119']

EXCLUDE_FOLDERS = {'client_selected_results'}

DETECTION_TYPE_MAP = {
    'proximity_detection_results': 'Proximity Violation',
    'phone_detection_results': 'Phone Detection',
    'crowd_detection_results': 'Crowd Detection',
    'entry_exit_detection_results': 'Entry/Exit Detection',
    'server_room_detection_results': 'Server Room Detection',
    'intruder_detection_results': 'Intruder Detection',
    'invigilator_detection_results': 'Invigilator Monitoring',
    'tcs_detections': 'TCS Detection',
}

print("="*80)
print("CGLE TIER 2 EXAM REPORT GENERATOR")
print("Jan 19, 2026 (Paper I & Paper II)")
print("="*80)

# Load Camera Metadata - Combine both sheets
print("\n📊 Step 1: Loading camera metadata...")

# Load Paper I sheet
df_paper1 = pd.read_excel(CAMERA_SHEET_PAPER1)
if 'Devie ID' in df_paper1.columns:
    df_paper1.rename(columns={'Devie ID': 'Device ID'}, inplace=True)

# Load Paper II sheet
df_paper2 = pd.read_excel(CAMERA_SHEET_PAPER2)
if 'Devie ID' in df_paper2.columns:
    df_paper2.rename(columns={'Devie ID': 'Device ID'}, inplace=True)

# Combine both sheets
df_cameras = pd.concat([df_paper1, df_paper2], ignore_index=True)
df_cameras = df_cameras.drop_duplicates(subset=['Device ID', 'Camera No'], keep='first')

df_cameras['Device ID'] = df_cameras['Device ID'].astype(str).str.strip()
df_cameras['Camera No'] = df_cameras['Camera No'].astype(int)
print(f"✓ Loaded {len(df_paper1)} cameras from Paper I sheet")
print(f"✓ Loaded {len(df_paper2)} cameras from Paper II sheet")
print(f"✓ Total unique cameras: {len(df_cameras)}")

cameras = {}
camera_name_lookup = {}

for _, row in df_cameras.iterrows():
    key = f"{row['Device ID']}_{row['Camera No']}"
    cameras[key] = {
        'Device ID': row['Device ID'],
        'Camera No': row['Camera No'],
        'Camera Name': row.get('Camera Name', ''),
        'Center Name': row.get('Center Name', ''),
        'City': row.get('City', ''),
        'TC Code': row.get('TC Code', ''),
        'State': row.get('State', ''),
    }
    
    name = str(row.get('Camera Name', '')).upper().strip()
    if name:
        if name not in camera_name_lookup:
            camera_name_lookup[name] = []
        camera_name_lookup[name].append(key)

# Helper Functions
def classify_camera_type(camera_name):
    if not camera_name:
        return "Outside Area"
    name_upper = camera_name.upper()
    if any(kw in name_upper for kw in ['SERVER', 'RACK', 'DATA CENTER']):
        return "Server Room"
    if any(kw in name_upper for kw in ['LAB', 'EXAM HALL', 'CLASS']) and re.search(r'\d', name_upper):
        return "Computer Lab"
    if any(kw in name_upper for kw in ['ENTRY', 'ENTRANCE', 'GATE', 'DOOR', 'MAIN', 'RECEPTION', 'FRISKING']):
        return "Entry"
    if any(kw in name_upper for kw in ['PARKING', 'VEHICLE']):
        return "Parking Lot"
    if any(kw in name_upper for kw in ['OUTDOOR', 'GROUND', 'GARDEN']):
        return "Outdoor Scene"
    return "Outside Area"

def determine_shift(hour, minute):
    """CGLE Tier 2 timing: 2 shifts with different durations"""
    time_in_minutes = hour * 60 + minute
    if 11*60 <= time_in_minutes < 13*60+15:     # 11:00 AM - 1:15 PM (Paper I)
        return "Shift 1"
    elif 16*60+30 <= time_in_minutes < 18*60+30: # 4:30 PM - 6:30 PM (Paper II)
        return "Shift 2"
    return None

# Scan Detection Files
print("\n🔍 Step 2: Scanning detection files...")

detections = defaultdict(lambda: defaultdict(list))
camera_counts = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
detection_files = defaultdict(lambda: defaultdict(list))

total_files_scanned = 0
total_files_matched = 0

pattern_violation = re.compile(r'violation_(\d+)_(\d+)_(\d{8})_(\d{6})')
pattern_phone = re.compile(r'phone_detected_(\d+)_(\d+)_(\d{8})_(\d{6})')
pattern_persons = re.compile(r'^(\d+)_(\d+)_(\d{8})_(\d{6})_')
pattern_no_targets = re.compile(r'(.+)_no_targets_(\d{8})_(\d{6})')
pattern_datetime = re.compile(r'(\d{8})_(\d{6})')
date_filter = re.compile(r'20260119')

for root, dirs, files in os.walk(RESULTS_DIR):
    folder_name = os.path.basename(root)
    if folder_name in EXCLUDE_FOLDERS:
        continue
    
    detection_type = DETECTION_TYPE_MAP.get(folder_name, folder_name.replace('_', ' ').title())
    
    for file in files:
        if not file.lower().endswith(('.jpg', '.png', '.jpeg')):
            continue
        
        if not date_filter.search(file):
            continue
        
        total_files_scanned += 1
        
        device_id, camera_no, date_str, time_str, camera_name = None, None, None, None, None
        
        if (m := pattern_violation.search(file)):
            device_id, camera_no, date_str, time_str = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        elif (m := pattern_phone.search(file)):
            device_id, camera_no, date_str, time_str = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        elif (m := pattern_persons.search(file)):
            device_id, camera_no, date_str, time_str = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        elif (m := pattern_no_targets.search(file)):
            camera_name = m.group(1).upper().strip()
            date_str, time_str = m.group(2), m.group(3)
        elif (m := pattern_datetime.search(file)):
            date_str, time_str = m.group(1), m.group(2)
        
        if not date_str or date_str not in DATES:
            continue
        
        if camera_name:
            matched_cam_id = None
            if camera_name in camera_name_lookup:
                matched_cam_id = camera_name_lookup[camera_name][0]
            else:
                for name, cam_ids in camera_name_lookup.items():
                    if camera_name in name or name in camera_name:
                        matched_cam_id = cam_ids[0]
                        break
            if matched_cam_id:
                device_id, camera_no = matched_cam_id.split('_')
                camera_no = int(camera_no)
        
        if not device_id or not date_str:
            continue
        
        try:
            dt = datetime.strptime(date_str + time_str, '%Y%m%d%H%M%S')
        except:
            continue
        
        shift = determine_shift(dt.hour, dt.minute)
        if not shift:
            continue
        
        cam_id = f"{device_id}_{camera_no}"
        cam_info = cameras.get(cam_id, {})
        if not cam_info:
            continue
        
        file_path = os.path.join(root, file).replace(BASE_DIR + '/', '')
        file_full_path = os.path.join(root, file)
        image_url = IMAGE_BASE_URL + file_path
        
        detection = {
            'Shift': shift,
            'Device ID': device_id,
            'Camera No': camera_no,
            'Cam ID': cam_id,
            'Camera Name': cam_info.get('Camera Name', ''),
            'Center Name': cam_info.get('Center Name', ''),
            'City': cam_info.get('City', ''),
            'Detection Type': detection_type,
            'Detection Time': dt,
            'File Path': file_path,
            'Image URL': image_url,
        }
        
        detections[date_str][shift].append(detection)
        camera_counts[date_str][shift][cam_id] += 1
        detection_files[date_str][shift].append(file_full_path)
        total_files_matched += 1

print(f"✓ Scanned {total_files_scanned} files")
print(f"✓ Matched {total_files_matched} detections")

# Create Proof of Work Folders
print("\n📁 Step 3: Creating proof of work folders...")

proof_of_work_dir = os.path.join(OUTPUT_DIR, 'proof_of_work')
os.makedirs(proof_of_work_dir, exist_ok=True)

total_copied = 0
for date_str in DATES:
    date_formatted = datetime.strptime(date_str, '%Y%m%d').strftime('%Y-%m-%d')
    
    for shift in ['Shift 1', 'Shift 2']:
        files = detection_files.get(date_str, {}).get(shift, [])
        if files:
            shift_name = shift.lower().replace(' ', '_')
            shift_dir = os.path.join(proof_of_work_dir, date_formatted, shift_name)
            os.makedirs(shift_dir, exist_ok=True)
            
            for file_path in files:
                if os.path.exists(file_path):
                    dest_path = os.path.join(shift_dir, os.path.basename(file_path))
                    shutil.copy2(file_path, dest_path)
                    total_copied += 1
            print(f"  ✓ {date_formatted}/{shift_name}: {len(files)} images")

print(f"✓ Total images copied: {total_copied}")

# Generate Excel Files
print("\n📝 Step 4: Generating Excel file...")

os.makedirs(OUTPUT_DIR, exist_ok=True)

def create_shift_data(cameras_dict, counts_dict, date_str, shift):
    data = []
    for cam_id, cam_info in sorted(cameras_dict.items()):
        event_count = counts_dict.get(date_str, {}).get(shift, {}).get(cam_id, 0)
        camera_status = "Online" if event_count > 0 else "Offline"
        cam_type = classify_camera_type(cam_info.get('Camera Name', ''))
        
        data.append({
            'City': cam_info.get('City', ''),
            'Center': cam_info.get('Center Name', ''),
            'Cam ID': cam_id,
            'Camera Name': cam_info.get('Camera Name', ''),
            'Cam Type': cam_type,
            'TC Code': cam_info.get('TC Code', ''),
            'Events(Excluding Rejected Events)': event_count,
            'Rejected Events': 0,
            'Camera Status': camera_status,
        })
    return pd.DataFrame(data)

def create_analysis_data(cameras_dict, counts_dict, detections_dict, date_str):
    date_formatted = datetime.strptime(date_str, '%Y%m%d').strftime('%B %d, %Y')
    total_cameras = len(cameras_dict)
    
    shift1_counts = counts_dict.get(date_str, {}).get('Shift 1', {})
    shift1_online = len([c for c, count in shift1_counts.items() if count > 0])
    shift1_total_detections = sum(shift1_counts.values())
    
    shift2_counts = counts_dict.get(date_str, {}).get('Shift 2', {})
    shift2_online = len([c for c, count in shift2_counts.items() if count > 0])
    shift2_total_detections = sum(shift2_counts.values())
    
    all_online = set(list(shift1_counts.keys()) + list(shift2_counts.keys()))
    unique_online = len(all_online)
    both_shifts = len(set(shift1_counts.keys()) & set(shift2_counts.keys()))
    only_shift1 = len(set(shift1_counts.keys()) - set(shift2_counts.keys()))
    only_shift2 = len(set(shift2_counts.keys()) - set(shift1_counts.keys()))
    
    def get_top_cameras(counts, n=5):
        sorted_cams = sorted(counts.items(), key=lambda x: x[1], reverse=True)[:n]
        return [(cid, cnt, cameras_dict[cid].get('Camera Name', '')) for cid, cnt in sorted_cams]
    
    shift1_top = get_top_cameras(shift1_counts)
    shift2_top = get_top_cameras(shift2_counts)
    
    data = []
    data.append(['DAILY CAMERA ANALYSIS REPORT', ''])
    data.append(['Date', date_formatted])
    data.append(['Exam', 'CGLE Tier 2'])
    data.append(['', ''])
    
    # Shift 1 - Paper I
    data.append(['SHIFT 1 ANALYSIS - PAPER I', ''])
    data.append(['Shift Timing', '11:00 AM - 1:15 PM'])
    data.append(['Duration', '2 hours 15 minutes'])
    data.append(['Total Cameras Monitored', total_cameras])
    data.append(['Online Cameras', shift1_online])
    data.append(['Offline Cameras', total_cameras - shift1_online])
    data.append(['Total Detections', shift1_total_detections])
    data.append(['Online Rate', f"{(shift1_online/total_cameras*100):.2f}%" if total_cameras > 0 else "0%"])
    data.append(['', ''])
    data.append(['Top 5 Most Active Cameras (Shift 1)', ''])
    for i, (cid, cnt, name) in enumerate(shift1_top, 1):
        data.append([f"{i}. {cid} ({name})", f"{cnt} detections"])
    if not shift1_top:
        data.append(['No active cameras', ''])
    data.append(['', ''])
    
    # Shift 2 - Paper II
    data.append(['SHIFT 2 ANALYSIS - PAPER II (STATISTICS)', ''])
    data.append(['Shift Timing', '4:30 PM - 6:30 PM'])
    data.append(['Duration', '2 hours'])
    data.append(['Total Cameras Monitored', total_cameras])
    data.append(['Online Cameras', shift2_online])
    data.append(['Offline Cameras', total_cameras - shift2_online])
    data.append(['Total Detections', shift2_total_detections])
    data.append(['Online Rate', f"{(shift2_online/total_cameras*100):.2f}%" if total_cameras > 0 else "0%"])
    data.append(['', ''])
    data.append(['Top 5 Most Active Cameras (Shift 2)', ''])
    for i, (cid, cnt, name) in enumerate(shift2_top, 1):
        data.append([f"{i}. {cid} ({name})", f"{cnt} detections"])
    if not shift2_top:
        data.append(['No active cameras', ''])
    data.append(['', ''])
    
    # Overall
    data.append(['OVERALL SUMMARY', ''])
    data.append(['Total Unique Cameras Online Today', unique_online])
    data.append(['Cameras Online in Both Shifts', both_shifts])
    data.append(['Cameras Online Only in Shift 1 (Paper I)', only_shift1])
    data.append(['Cameras Online Only in Shift 2 (Paper II)', only_shift2])
    data.append(['Total Detections (Both Papers)', shift1_total_detections + shift2_total_detections])
    data.append(['Total Exam Duration', '4 hours 15 minutes (2 papers)'])
    
    return pd.DataFrame(data, columns=['Metric', 'Value'])

def create_detections_list(detections_dict, date_str):
    all_detections = []
    for shift in ['Shift 1', 'Shift 2']:
        all_detections.extend(detections_dict.get(date_str, {}).get(shift, []))
    all_detections.sort(key=lambda x: x['Detection Time'])
    for det in all_detections:
        det['Detection Time'] = det['Detection Time'].strftime('%Y-%m-%d %H:%M:%S')
        det.pop('File Path', None)
    return pd.DataFrame(all_detections)

# Generate single file for the date
date_str = DATES[0]
date_obj = datetime.strptime(date_str, '%Y%m%d')
file_date = date_obj.strftime('%Y-%m-%d')
output_file = os.path.join(OUTPUT_DIR, f"{file_date}.xlsx")

print(f"  Creating {file_date}.xlsx...")

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    create_shift_data(cameras, camera_counts, date_str, 'Shift 1').to_excel(writer, sheet_name='Shift 1 - Paper I', index=False)
    create_shift_data(cameras, camera_counts, date_str, 'Shift 2').to_excel(writer, sheet_name='Shift 2 - Paper II', index=False)
    create_analysis_data(cameras, camera_counts, detections, date_str).to_excel(writer, sheet_name='Analysis', index=False)
    
    df_det = create_detections_list(detections, date_str)
    if not df_det.empty:
        df_det.to_excel(writer, sheet_name='Detections List', index=False)
    else:
        pd.DataFrame(columns=['Shift', 'Device ID', 'Camera No', 'Cam ID', 'Camera Name', 
                             'Center Name', 'City', 'Detection Type', 'Detection Time', 
                             'Image URL']).to_excel(writer, sheet_name='Detections List', index=False)

wb = load_workbook(output_file)
if 'Detections List' in wb.sheetnames:
    ws = wb['Detections List']
    max_col = ws.max_column
    ws.cell(row=1, column=max_col + 1, value='View').font = Font(bold=True)
    for row_idx in range(2, ws.max_row + 1):
        url = ws.cell(row=row_idx, column=max_col).value
        if url:
            cell = ws.cell(row=row_idx, column=max_col + 1)
            cell.value = 'View'
            cell.hyperlink = url
            cell.font = Font(color="0000FF", underline="single")

if 'Analysis' in wb.sheetnames:
    ws = wb['Analysis']
    ws.column_dimensions['A'].width = 50
    ws.column_dimensions['B'].width = 30
    for cell in ws['A']:
        if cell.value and any(kw in str(cell.value) for kw in ['ANALYSIS', 'SUMMARY', 'REPORT']):
            cell.font = Font(bold=True, size=12)

wb.save(output_file)

shift1_count = sum(camera_counts.get(date_str, {}).get('Shift 1', {}).values())
shift2_count = sum(camera_counts.get(date_str, {}).get('Shift 2', {}).values())

print(f"    ✓ Shift 1 (Paper I): {shift1_count} events")
print(f"    ✓ Shift 2 (Paper II): {shift2_count} events")

print("\n" + "="*80)
print("✅ CGLE TIER 2 EXAM REPORTS COMPLETE")
print("="*80)
print(f"Output: {OUTPUT_DIR}")
print(f"  - 1 Excel file with 4 sheets")
print(f"  - proof_of_work/ folder with detection images")
