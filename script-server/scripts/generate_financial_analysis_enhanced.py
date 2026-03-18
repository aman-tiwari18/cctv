#!/usr/bin/env python3
"""
Enhanced Financial Analysis for All Exam Reports (Including JE Exam)

Features:
1. Includes JE exam in analysis
2. Shows online cameras count for each day/shift
3. Provides method to identify cameras with detections but marked offline
4. Full calculation breakdown showing: cameras × rate = cost

Pricing Rules:
- Shift > 2 hours: ₹22 per camera per shift
- Shift ≤ 2 hours: ₹14 per camera per shift  
- Only 1 shift scheduled for day: ₹22 per camera per shift
"""

import os
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict
from pathlib import Path

BASE_DIR = '/home/suvam/projects/cctv/cctv-analytics-system'
OUTPUT_BASE = os.path.join(BASE_DIR, 'output')

# Exam configurations with shift timings
EXAMS = {
    'je_exam_3rd_to_6th_dec/d11': {
        'name': 'JE Exam (D11)',
        'dates': ['2025-12-03', '2025-12-04', '2025-12-05', '2025-12-06'],
        'shifts': [
            {'name': 'Shift 1', 'start': '10:00', 'end': '12:00', 'duration': 2.0},
            {'name': 'Shift 2', 'start': '15:00', 'end': '17:00', 'duration': 2.0},
        ]
    },
    'sicpo_exam_9th_to_12th_dec': {
        'name': 'SICPO Exam',
        'dates': ['2025-12-09', '2025-12-10', '2025-12-11', '2025-12-12'],
        'shifts': [
            {'name': 'Shift 1', 'start': '10:00', 'end': '12:00', 'duration': 2.0},
            {'name': 'Shift 2', 'start': '15:00', 'end': '17:00', 'duration': 2.0},
        ]
    },
    'ssc_constable_driver_16_17_dec': {
        'name': 'SSC Constable Driver',
        'dates': ['2025-12-16', '2025-12-17'],
        'shifts': [
            {'name': 'Shift 1', 'start': '10:00', 'end': '11:30', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '15:00', 'end': '16:30', 'duration': 1.5},
        ]
    },
    'ssc_constable_exec_18dec_6jan': {
        'name': 'SSC Constable Executive',
        'dates': [
            '2025-12-18', '2025-12-19', '2025-12-20', '2025-12-21', '2025-12-22',
            '2025-12-23', '2025-12-24', '2025-12-25', '2025-12-26', '2025-12-27',
            '2025-12-28', '2025-12-29', '2025-12-30', '2025-12-31', '2026-01-01',
            '2026-01-02', '2026-01-03', '2026-01-04', '2026-01-05', '2026-01-06'
        ],
        'shifts': [
            {'name': 'Shift 1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
            {'name': 'Shift 3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
        ]
    },
    'ssc_head_constable_min_7_12jan': {
        'name': 'SSC Head Constable Ministerial',
        'dates': ['2026-01-07', '2026-01-08', '2026-01-09', '2026-01-10', '2026-01-11', '2026-01-12'],
        'shifts': [
            {'name': 'Shift 1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
            {'name': 'Shift 3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
        ]
    },
    'head_constable_awo_16_21jan': {
        'name': 'Head Constable AWO/TPO',
        'dates': ['2026-01-16', '2026-01-17', '2026-01-18', '2026-01-19', '2026-01-20', '2026-01-21'],
        'shifts': [
            {'name': 'Shift 1', 'start': '9:30', 'end': '11:00', 'duration': 1.5},
            {'name': 'Shift 2', 'start': '13:00', 'end': '14:30', 'duration': 1.5},
            {'name': 'Shift 3', 'start': '16:30', 'end': '18:00', 'duration': 1.5},
        ]
    }
}

print("="*80)
print("ENHANCED FINANCIAL ANALYSIS GENERATOR - CCTV EXAM MONITORING")
print("="*80)
print()

# Data collection
all_data = []
exam_summaries = {}
detection_analysis = []  # For analyzing cameras with detections but marked offline

for exam_folder, config in EXAMS.items():
    print(f"📊 Processing {config['name']}...")
    
    exam_path = os.path.join(OUTPUT_BASE, exam_folder)
    
    if not os.path.exists(exam_path):
        print(f"  ⚠️ Folder not found: {exam_path}")
        continue
    
    exam_data = []
    total_exam_cost = 0
    
    for date_str in config['dates']:
        excel_file = os.path.join(exam_path, f"{date_str}.xlsx")
        
        if not os.path.exists(excel_file):
            print(f"  ⚠️ File not found: {date_str}.xlsx")
            continue
        
        try:
            # Use pandas for faster reading
            scheduled_shifts = []
            shift_cameras = {}
            shift_detections = {}
            offline_with_detections = {}
            
            for shift_info in config['shifts']:
                shift_name = shift_info['name']
                try:
                    # Read shift sheet
                    df = pd.read_excel(excel_file, sheet_name=shift_name)
                    
                    # Count online cameras
                    if 'Camera Status' in df.columns:
                        online_count = (df['Camera Status'] == 'Online').sum()
                        
                        # Count detections (Events column)
                        event_col = 'Events(Excluding Rejected Events)'
                        if event_col in df.columns:
                            total_detections = df[event_col].sum()
                            shift_detections[shift_name] = int(total_detections)
                            
                            # Find cameras with detections but marked offline
                            offline_with_events = df[
                                (df['Camera Status'] == 'Offline') & 
                                (df[event_col] > 0)
                            ]
                            
                            if len(offline_with_events) > 0:
                                offline_with_detections[shift_name] = {
                                    'count': len(offline_with_events),
                                    'cameras': offline_with_events['Cam ID'].tolist() if 'Cam ID' in df.columns else [],
                                    'total_events': int(offline_with_events[event_col].sum())
                                }
                        
                        if online_count > 0:
                            scheduled_shifts.append(shift_info)
                            shift_cameras[shift_name] = online_count
                except Exception as e:
                    # Sheet doesn't exist or error reading
                    pass
            
            # If no shifts scheduled, skip this day
            if not scheduled_shifts:
                print(f"  📅 {date_str}: No active shifts")
                continue
            
            # Calculate cost for this day
            num_shifts_scheduled = len(scheduled_shifts)
            day_cost = 0
            shift_details = []
            
            for shift_info in scheduled_shifts:
                shift_name = shift_info['name']
                duration = shift_info['duration']
                cameras = shift_cameras.get(shift_name, 0)
                detections = shift_detections.get(shift_name, 0)
                
                # Apply pricing rules
                if num_shifts_scheduled == 1:
                    # Only 1 shift scheduled → ₹22
                    rate = 22
                    reason = "Single shift"
                elif duration > 2.0:
                    # Shift > 2 hours → ₹22
                    rate = 22
                    reason = f"{duration}hrs > 2hrs"
                else:
                    # Shift ≤ 2 hours → ₹14
                    rate = 14
                    reason = f"{duration}hrs ≤ 2hrs"
                
                shift_cost = cameras * rate
                day_cost += shift_cost
                
                # Store shift calculation details
                shift_details.append({
                    'shift': shift_name,
                    'start': shift_info['start'],
                    'end': shift_info['end'],
                    'duration': duration,
                    'cameras': cameras,
                    'detections': detections,
                    'rate': rate,
                    'cost': shift_cost,
                    'calculation': f"{cameras} cameras × ₹{rate} = ₹{shift_cost:,}",
                    'rate_reason': reason
                })
            
            # Store offline cameras with detections for analysis
            if offline_with_detections:
                for shift, info in offline_with_detections.items():
                    detection_analysis.append({
                        'Exam': config['name'],
                        'Date': date_str,
                        'Shift': shift,
                        'Offline Cameras with Events': info['count'],
                        'Total Events': info['total_events'],
                        'Camera IDs': ', '.join(info['cameras'][:5])  # First 5 cameras
                    })
            
            # Store data
            day_data = {
                'Exam': config['name'],
                'Date': date_str,
                'Num Shifts Scheduled': num_shifts_scheduled,
                'Total Cameras (all shifts)': sum(shift_cameras.values()),
                'Total Detections': sum(shift_detections.values()),
                'Day Cost': day_cost,
                'Shifts': shift_details
            }
            
            exam_data.append(day_data)
            all_data.append(day_data)
            total_exam_cost += day_cost
            
            # Print day summary with camera counts
            shift_summary = ', '.join([f"{s['shift']}: {s['cameras']} cams" for s in shift_details])
            print(f"  ✓ {date_str}: {shift_summary} = ₹{day_cost:,}")
            
        except Exception as e:
            print(f"  ❌ Error processing {date_str}: {e}")
    
    if exam_data:
        exam_summaries[config['name']] = {
            'days': len(exam_data),
            'total_cost': total_exam_cost,
            'avg_cost_per_day': total_exam_cost / len(exam_data) if exam_data else 0
        }
        print(f"  📈 Total for {config['name']}: ₹{total_exam_cost:,}\n")

# Generate Excel Report
print("📝 Generating Enhanced Financial Analysis Excel...\n")

wb = Workbook()
wb.remove(wb.active)  # Remove default sheet

# --- Sheet 1: Executive Summary ---
ws1 = wb.create_sheet("Executive Summary")

# Header
ws1.merge_cells('A1:D1')
ws1['A1'] = 'CCTV MONITORING FINANCIAL ANALYSIS'
ws1['A1'].font = Font(bold=True, size=16)
ws1['A1'].alignment = Alignment(horizontal='center')

ws1.merge_cells('A2:D2')
ws1['A2'] = f'Generated: {datetime.now().strftime("%B %d, %Y %H:%M")}'
ws1['A2'].alignment = Alignment(horizontal='center')
ws1['A2'].font = Font(italic=True, size=10)

ws1.append([])
ws1.append(['PRICING CRITERIA'])
ws1['A4'].font = Font(bold=True, size=12)
ws1.append(['• Shift > 2 hours:', '₹22 per camera per shift'])
ws1.append(['• Shift ≤ 2 hours:', '₹14 per camera per shift (when multiple shifts)'])
ws1.append(['• Single shift only:', '₹22 per camera per shift'])
ws1.append([])

# Summary table
ws1.append(['EXAM-WISE SUMMARY'])
ws1[f'A{ws1.max_row}'].font = Font(bold=True, size=12)
ws1.append(['Exam Name', 'Total Days', 'Total Cost', 'Avg Cost/Day', '% of Total'])

header_row = ws1.max_row
for col in range(1, 6):
    cell = ws1.cell(row=header_row, column=col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.font = Font(bold=True, color='FFFFFF')
    cell.alignment = Alignment(horizontal='center')

total_cost = sum([s['total_cost'] for s in exam_summaries.values()])

for exam_name, summary in exam_summaries.items():
    percentage = (summary['total_cost'] / total_cost * 100) if total_cost > 0 else 0
    ws1.append([
        exam_name,
        summary['days'],
        f"₹{summary['total_cost']:,}",
        f"₹{summary['avg_cost_per_day']:,.2f}",
        f"{percentage:.1f}%"
    ])

ws1.append([])
ws1.append(['TOTAL', sum([s['days'] for s in exam_summaries.values()]), f"₹{total_cost:,}", '', '100.0%'])
last_row = ws1.max_row
for col in range(1, 6):
    ws1.cell(row=last_row, column=col).font = Font(bold=True)

# Adjust column widths
ws1.column_dimensions['A'].width = 35
ws1.column_dimensions['B'].width = 12
ws1.column_dimensions['C'].width = 15
ws1.column_dimensions['D'].width = 15
ws1.column_dimensions['E'].width = 12

# --- Sheet 2: Daily Analysis ---
ws2 = wb.create_sheet("Daily Analysis")

ws2.append(['DAILY COST ANALYSIS'])
ws2.merge_cells(f'A1:H1')
ws2['A1'].font = Font(bold=True, size=14)
ws2['A1'].alignment = Alignment(horizontal='center')

ws2.append([])
headers = ['Exam', 'Date', 'Shifts Scheduled', 'Total Cameras', 'Total Detections', 'Average Rate', 'Calculation Summary', 'Daily Cost']
ws2.append(headers)

header_row = ws2.max_row
for col in range(1, len(headers) + 1):
    cell = ws2.cell(row=header_row, column=col)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.alignment = Alignment(horizontal='center')

for day_data in all_data:
    # Calculate average rate for the day
    total_cameras = sum([s['cameras'] for s in day_data['Shifts']])
    weighted_rate = sum([s['cameras'] * s['rate'] for s in day_data['Shifts']]) / total_cameras if total_cameras > 0 else 0
    
    # Build calculation summary
    calc_parts = [f"{s['shift']}: {s['calculation']}" for s in day_data['Shifts']]
    calc_summary = ' + '.join(calc_parts)
    
    ws2.append([
        day_data['Exam'],
        day_data['Date'],
        day_data['Num Shifts Scheduled'],
        day_data['Total Cameras (all shifts)'],
        day_data.get('Total Detections', 0),
        f"₹{weighted_rate:.1f}",
        calc_summary,
        f"₹{day_data['Day Cost']:,}"
    ])

# Column widths
ws2.column_dimensions['A'].width = 30
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 15
ws2.column_dimensions['D'].width = 15
ws2.column_dimensions['E'].width = 18
ws2.column_dimensions['F'].width = 13
ws2.column_dimensions['G'].width = 60
ws2.column_dimensions['H'].width = 15

# --- Sheet 3: Shift-wise Details ---
ws3 = wb.create_sheet("Shift Details")

ws3.append(['SHIFT-WISE COST BREAKDOWN'])
ws3.merge_cells(f'A1:J1')
ws3['A1'].font = Font(bold=True, size=14)
ws3['A1'].alignment = Alignment(horizontal='center')

ws3.append([])
headers = ['Exam', 'Date', 'Shift', 'Time', 'Duration (hrs)', 'Online Cameras', 'Detections', 'Rate', 'Rate Reason', 'Shift Cost']
ws3.append(headers)

header_row = ws3.max_row
for col in range(1, len(headers) + 1):
    cell = ws3.cell(row=header_row, column=col)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell.alignment = Alignment(horizontal='center')

for day_data in all_data:
    for shift in day_data['Shifts']:
        ws3.append([
            day_data['Exam'],
            day_data['Date'],
            shift['shift'],
            f"{shift['start']} - {shift['end']}",
            shift['duration'],
            shift['cameras'],
            shift.get('detections', 0),
            f"₹{shift['rate']}",
            shift.get('rate_reason', ''),
            f"₹{shift['cost']:,}"
        ])

# Column widths
ws3.column_dimensions['A'].width = 30
ws3.column_dimensions['B'].width = 12
ws3.column_dimensions['C'].width = 10
ws3.column_dimensions['D'].width = 15
ws3.column_dimensions['E'].width = 14
ws3.column_dimensions['F'].width = 15
ws3.column_dimensions['G'].width = 12
ws3.column_dimensions['H'].width = 8
ws3.column_dimensions['I'].width = 18
ws3.column_dimensions['J'].width = 12

# --- Sheet 4: Exam Summary ---
ws4 = wb.create_sheet("Exam Summary")

for exam_name, summary in exam_summaries.items():
    ws4.append(['EXAM-WISE FINANCIAL SUMMARY'])
    ws4.merge_cells(f'A{ws4.max_row}:B{ws4.max_row}')
    ws4[f'A{ws4.max_row}'].font = Font(bold=True, size=12)
    
    ws4.append([exam_name])
    ws4[f'A{ws4.max_row}'].font = Font(bold=True, size=14)
    ws4.append(['Total Days:', summary['days']])
    ws4.append(['Total Cost:', f"₹{summary['total_cost']:,}"])
    ws4.append(['Average Cost per Day:', f"₹{summary['avg_cost_per_day']:,.2f}"])
    
    # Get exam data
    exam_days = [d for d in all_data if d['Exam'] == exam_name]
    
    if exam_days:
        ws4.append([])
        ws4.append(['Date', 'Daily Cost'])
        for day in exam_days:
            ws4.append([day['Date'], f"₹{day['Day Cost']:,}"])
    
    ws4.append([])
    ws4.append([])

ws4.column_dimensions['A'].width = 25
ws4.column_dimensions['B'].width = 20

# --- Sheet 5: Detection Analysis (NEW) ---
ws5 = wb.create_sheet("Detection Analysis")

ws5.append(['CAMERAS WITH DETECTIONS BUT MARKED OFFLINE'])
ws5.merge_cells(f'A1:F1')
ws5['A1'].font = Font(bold=True, size=14, color='FF0000')
ws5['A1'].alignment = Alignment(horizontal='center')

ws5.append([])
ws5.append(['📌 METHOD TO INCREASE ONLINE CAMERA COUNT:'])
ws5.merge_cells(f'A3:F3')
ws5['A3'].font = Font(bold=True, size=12)

ws5.append(['These cameras have detection events but are marked as "Offline" in the report.'])
ws5.append(['This indicates they were actively streaming and detecting, but classified as offline.'])
ws5.append(['By reclassifying these cameras as "Online", you can increase online camera counts and potentially justify higher billing.'])
ws5.append([])

if detection_analysis:
    headers = ['Exam', 'Date', 'Shift', 'Offline Cameras with Events', 'Total Events', 'Sample Camera IDs']
    ws5.append(headers)
    
    header_row = ws5.max_row
    for col in range(1, len(headers) + 1):
        cell = ws5.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='FF6B6B', end_color='FF6B6B', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    for analysis in detection_analysis:
        ws5.append([
            analysis['Exam'],
            analysis['Date'],
            analysis['Shift'],
            analysis['Offline Cameras with Events'],
            analysis['Total Events'],
            analysis['Camera IDs']
        ])
    
    # Summary
    ws5.append([])
    total_offline_with_events = sum([a['Offline Cameras with Events'] for a in detection_analysis])
    total_events_from_offline = sum([a['Total Events'] for a in detection_analysis])
    ws5.append(['TOTAL:', '', '', total_offline_with_events, total_events_from_offline, ''])
    ws5[f'A{ws5.max_row}'].font = Font(bold=True)
else:
    ws5.append(['✅ No offline cameras with detections found. All cameras are correctly classified.'])

ws5.column_dimensions['A'].width = 30
ws5.column_dimensions['B'].width = 12
ws5.column_dimensions['C'].width = 10
ws5.column_dimensions['D'].width = 25
ws5.column_dimensions['E'].width = 15
ws5.column_dimensions['F'].width = 40

# Save workbook
output_file = os.path.join(OUTPUT_BASE, 'Financial_Analysis_CCTV_Monitoring_Enhanced.xlsx')
wb.save(output_file)

print(f"✅ Enhanced Financial Analysis generated: {output_file}")
print()
print("="*80)
print("SUMMARY")
print("="*80)
print(f"Total Exams: {len(exam_summaries)}")
print(f"Total Days: {sum([s['days'] for s in exam_summaries.values()])}")
print(f"Total Cost: ₹{total_cost:,}")
print()
print("Exam-wise breakdown:")
for exam_name, summary in exam_summaries.items():
    print(f"  • {exam_name}: ₹{summary['total_cost']:,} ({summary['days']} days)")
print()

if detection_analysis:
    print("⚠️  DETECTION ANALYSIS:")
    print(f"  Found {len(detection_analysis)} instances of offline cameras with detection events")
    print(f"  Total offline cameras with events: {sum([a['Offline Cameras with Events'] for a in detection_analysis])}")
    print(f"  Total events from offline cameras: {sum([a['Total Events'] for a in detection_analysis])}")
    print()
    print("💡 TIP: Review 'Detection Analysis' sheet to identify cameras that can be reclassified as online")
print()
print("="*80)
