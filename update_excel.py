import openpyxl

file_path = "/data/cctv_dev/cctv/CCTV infra and compute time.xlsx"
wb = openpyxl.load_workbook(file_path)
sheet = wb.active

updates = {
    10: {
        "cdis": "Real-time if manual operator push is removed, else ~5-10 seconds latency.",
        "verdict": "✓ Correct"
    }
}

for row in range(2, 12):
    stage_id = sheet.cell(row=row, column=1).value
    if stage_id in updates:
        sheet.cell(row=row, column=7).value = updates[stage_id]["cdis"]
        sheet.cell(row=row, column=8).value = updates[stage_id]["verdict"]

wb.save(file_path)
print("Updated Excel for Stage 10 successfully!")
